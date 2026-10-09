from __future__ import annotations

import asyncio
import csv
import difflib
import io
import json
import os
import re
import shutil
import tempfile
import threading
import time
import unicodedata
import uuid
import webbrowser
import zipfile
from collections import Counter, defaultdict
from copy import deepcopy
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from biblioref.projects import LEGACY_PROJECT_ID, PROJECTS_VERSION, project_tree, relative_path

from biblioref.ingestion.backend_context import pymupdf as fitz

from biblioref.audit.manual import process_manual_zone
from biblioref.audit.models import ManualZoneSpec
from biblioref.core.config import SOFTWARE_VERSION, ALGORITHM_VERSION, STABLE_ENGINE_VERSION
from biblioref.ingestion import (
    ZIP_INGESTION_VERSION,
    ZipIngestionError,
    extract_pdfs_from_zip,
    extract_layout_lines_range,
    page_dimensions,
    pdf_page_count,
    pdf_preflight,
    render_band_png,
    render_page_png,
)
from biblioref.pipeline import process_pdf
from biblioref.audit.coverage import build_coverage_report, cut_opens_a_record
from biblioref.parsing import PARSER_STATUS, PARSER_VERSION, flatten_parsed, parse_reference
from biblioref.parsing.inheritance import inherit_repeated_authors
from biblioref.review_problems import problem_counts
from biblioref.styleprofiler import STYLE_PROFILER_VERSION, profile_document as profile_style_document
from biblioref.integrations import google_drive as gdrive
from biblioref.metadata.inference import field_origin
from biblioref.duplicates import DuplicateDocument, content_hash, sha256_file
from biblioref.resolver.interactive import InteractiveResolutionRunner

try:
    import openpyxl
    from openpyxl import Workbook
except Exception:  # pragma: no cover
    openpyxl = None
    Workbook = None

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
WORKSPACE = Path(os.environ.get("BIBLIOREF_WORKSPACE", Path.home() / ".biblioref" / "projects" / "default"))
UPLOAD_DIR = WORKSPACE / "pdfs"
RESULT_DIR = WORKSPACE / "results"
STATE_FILE = WORKSPACE / "project.json"
for folder in (WORKSPACE, UPLOAD_DIR, RESULT_DIR):
    folder.mkdir(parents=True, exist_ok=True)

# APP-level extraction revision. This is intentionally separate from the
# algorithm version: it invalidates project snapshots created by earlier APP
# bundles that could claim v0.9.6 provenance while still carrying an older
# persisted reference list.
EXTRACTION_REVISION = "v096-core-first-reference-1"
APP_VERSION = "2.7.0"
PDF_CLICK_MARGIN_PT = 42.0
UI_WORKBENCH_VERSION = "ui-reference-workbench-v0.3-no-local-runtime"

PARSED_REVIEW_FIELDS = (
    "author", "year", "publication_type", "title", "journal",
    "publisher", "publisher_observed", "source_publisher", "doi",
)


def _now() -> int:
    return int(time.time())


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _join_spans_for_reading(spans: list[dict[str, Any]]) -> str:
    """Join PDF spans by physical gap instead of inserting a space blindly.

    This is a presentation/title helper only. The extractor keeps its original
    text path unchanged so regression and exported raw references stay stable.
    """
    parts = [sp for sp in spans if str(sp.get("text") or "").strip()]
    if not parts:
        return ""
    parts.sort(key=lambda sp: float((sp.get("bbox") or (0, 0, 0, 0))[0]))
    out = str(parts[0].get("text") or "").strip()
    prev = parts[0]
    for cur in parts[1:]:
        raw_prev = str(prev.get("text") or "")
        raw_cur = str(cur.get("text") or "")
        left = raw_prev.rstrip()
        right = raw_cur.lstrip()
        pb = prev.get("bbox") or (0, 0, 0, 0)
        cb = cur.get("bbox") or (0, 0, 0, 0)
        gap = float(cb[0]) - float(pb[2])
        size = max(1.0, (float(prev.get("size") or 0) + float(cur.get("size") or 0)) / 2.0)
        explicit_space = raw_prev.endswith((" ", "\t")) or raw_cur.startswith((" ", "\t"))
        # A normal printed inter-word space is usually clearly larger than the
        # tiny gap produced by a font/style span switch inside the same word.
        separator = " " if explicit_space or gap > max(1.15, size * 0.19) else ""
        out = out.rstrip() + separator + right
        prev = cur
    return _norm(out)


def _safe_name(name: str) -> str:
    name = Path(name).name
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .")
    return stem or "document.pdf"


def _filename_as_title(filename: str) -> str:
    text = Path(filename).stem
    text = re.sub(r"[_-]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or Path(filename).stem


_TITLE_STOP_TERMS = (
    "universidad", "university", "facultad", "faculty", "departamento", "department",
    "tesis doctoral", "doctoral thesis", "doctoral dissertation", "dissertation", "tesis de",
    "presentada por", "presented by", "autor:", "author:", "director", "supervisor",
    "doctorado", "doctoral program", "licenciatura", "maestría", "master thesis", "conicet",
)


def _plausible_title(text: str) -> bool:
    t = _norm(text).strip(" -–—:;,.|")
    if len(t) < 12 or len(t) > 320:
        return False
    low = t.lower()
    if low in {"tesis", "tesis doctoral", "doctoral thesis", "dissertation", "bibliografía", "referencias"}:
        return False
    words = re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñÀ-ÿ]+", t)
    return len(words) >= 3


def _detect_thesis_title(path: Path, filename: str) -> tuple[str, str]:
    """Best-effort title suggestion. It is always editable in the UI."""
    fallback = _filename_as_title(filename)
    if fitz is None:
        return fallback, "filename"
    try:
        with fitz.open(path) as pdf:
            meta = _norm((pdf.metadata or {}).get("title") or "")
            if _plausible_title(meta) and meta.lower() not in {fallback.lower(), filename.lower()}:
                if not any(term == meta.lower() for term in _TITLE_STOP_TERMS):
                    return meta, "pdf_metadata"

            groups = []
            for pno in range(min(3, len(pdf))):
                page = pdf[pno]
                pdata = page.get_text("dict")
                height = float(page.rect.height or 1)
                lines = []
                for block in pdata.get("blocks", []):
                    if block.get("type") != 0:
                        continue
                    for line in block.get("lines", []):
                        spans = line.get("spans", [])
                        text = _join_spans_for_reading(spans)
                        if not _plausible_title(text):
                            continue
                        size = max((float(sp.get("size") or 0) for sp in spans), default=0.0)
                        bbox = line.get("bbox") or (0, 0, 0, 0)
                        lines.append({"text": text, "size": size, "y0": float(bbox[1]), "y1": float(bbox[3]), "height": height})
                if not lines:
                    continue
                max_size = max(x["size"] for x in lines)
                threshold = max(11.0, max_size - 3.0)
                big = [x for x in lines if x["size"] >= threshold]
                big.sort(key=lambda x: x["y0"])
                current = []
                for ln in big:
                    if current and ln["y0"] - current[-1]["y1"] > max(45.0, 3.2 * max(ln["size"], current[-1]["size"])):
                        groups.append((pno, current))
                        current = []
                    current.append(ln)
                if current:
                    groups.append((pno, current))

            scored = []
            for pno, group in groups:
                text = _norm(" ".join(x["text"] for x in group))
                if not _plausible_title(text):
                    continue
                low = text.lower()
                size = max(x["size"] for x in group)
                ymid = sum((x["y0"] + x["y1"]) / 2 for x in group) / len(group)
                height = group[0]["height"]
                generic_hits = sum(term in low for term in _TITLE_STOP_TERMS)
                institution_only = any(k in low for k in ("universidad", "university", "facultad", "faculty")) and len(text.split()) < 14
                score = size * 2.0 + min(len(text), 180) / 18.0
                score += 5.0 if 0.18 * height <= ymid <= 0.72 * height else -2.0
                score -= generic_hits * 10.0
                score -= 12.0 if institution_only else 0.0
                score -= pno * 2.0
                scored.append((score, text))
            if scored:
                scored.sort(reverse=True, key=lambda x: x[0])
                best = scored[0][1].strip(" -–—:;,.|")
                if _plausible_title(best):
                    return best, "cover_layout"
    except Exception:
        pass
    return fallback, "filename"


def _next_thesis_id() -> str:
    nums = []
    for doc in STORE.docs().values() if 'STORE' in globals() else []:
        m = re.fullmatch(r"T(\d+)", str(doc.get("thesis_id") or ""))
        if m:
            nums.append(int(m.group(1)))
    return f"T{(max(nums, default=0) + 1):04d}"


TITLE_IDENTITY_REVISION = "gap-spans-1"


def _ensure_doc_identity(doc: dict[str, Any]) -> bool:
    changed = False
    if not doc.get("thesis_id"):
        doc["thesis_id"] = _next_thesis_id()
        changed = True
    if not isinstance(doc.get("metadata"), dict):
        doc["metadata"] = {}
        changed = True
    if not doc.get("title"):
        path = Path(doc.get("path") or "")
        if path.exists():
            title, source = _detect_thesis_title(path, doc.get("filename") or path.name)
        else:
            title, source = _filename_as_title(doc.get("filename") or "tesis"), "filename"
        doc["title"] = title
        doc["title_source"] = source
        doc["title_confirmed"] = False
        doc["title_reading_revision"] = TITLE_IDENTITY_REVISION
        changed = True
    elif not doc.get("title_confirmed") and doc.get("title_reading_revision") != TITLE_IDENTITY_REVISION:
        path = Path(doc.get("path") or "")
        if path.exists():
            title, source = _detect_thesis_title(path, doc.get("filename") or path.name)
            if title:
                doc["title"] = title
                doc["title_source"] = source
        doc["title_reading_revision"] = TITLE_IDENTITY_REVISION
        changed = True
    return changed


def _inferred_fields(doc: dict[str, Any]) -> str:
    """Thesis fields whose value was inferred rather than read from the PDF, for exports."""
    return "; ".join(f for f in ("author", "year", "degree_level", "program", "faculty", "university")
                     if field_origin(doc, f) == "inferido")


def _analysis_versions(doc: dict[str, Any]) -> dict[str, Any]:
    analysis = doc.get("analysis") or {}
    return {
        "software_version": analysis.get("software_version"),
        "algorithm_version": analysis.get("algorithm_version"),
        "stable_engine_version": analysis.get("stable_engine_version"),
        "extraction_revision": analysis.get("extraction_revision"),
    }


def _analysis_stale(doc: dict[str, Any]) -> bool:
    if doc.get("status") in {"uploaded", "processing"} or not doc.get("references"):
        return False
    v = _analysis_versions(doc)
    # Old projects may predate version provenance entirely; treat them as stale.
    return not (
        v.get("software_version") == SOFTWARE_VERSION
        and v.get("algorithm_version") == ALGORITHM_VERSION
        and v.get("stable_engine_version") == STABLE_ENGINE_VERSION
        and v.get("extraction_revision") == EXTRACTION_REVISION
    )


def _archive_analysis_snapshot(doc: dict[str, Any], reason: str = "reprocess") -> None:
    if not doc.get("references") and not doc.get("sections"):
        return
    snap = {
        "id": f"snap_{uuid.uuid4().hex[:10]}",
        "timestamp": _now(),
        "reason": reason,
        "versions": _analysis_versions(doc),
        "n_references": len([r for r in doc.get("references", []) if not r.get("deleted")]),
        "analysis": deepcopy(doc.get("analysis") or {}),
        "sections": deepcopy(doc.get("sections") or []),
        "references": deepcopy(doc.get("references") or []),
        "corrections": deepcopy(doc.get("corrections") or []),
    }
    doc.setdefault("analysis_history", []).append(snap)
    # Keep the project file bounded while preserving recent migrations.
    doc["analysis_history"] = doc["analysis_history"][-5:]


def _has_manual_work(doc: dict[str, Any]) -> bool:
    """Return True only for human work that would be lost by re-extraction.

    Thesis metadata edits and audit decisions are deliberately *not* blockers: they
    live outside the extracted reference list and survive a fresh extraction. Earlier
    APP bundles treated any correction log entry as manual extraction work, so merely
    correcting a title could freeze an old v0.9.5 reference list forever.
    """
    harmless_corrections = {"edit_thesis_metadata", "bulk_edit_thesis_metadata", "audit_check_decision"}
    for corr in doc.get("corrections") or []:
        kind = str(corr.get("kind") or "")
        if kind and kind not in harmless_corrections:
            return True
    for sec in doc.get("sections") or []:
        if sec.get("source") == "manual" or sec.get("confirmed"):
            return True
    for ref in doc.get("references") or []:
        if ref.get("deleted") or ref.get("edited") or ref.get("source") not in (None, "automatic"):
            return True
        if ref.get("status") not in (None, "extracted"):
            return True
    return False


def _public_doc(doc: dict[str, Any], include_refs: bool = False) -> dict[str, Any]:
    keep = {
        "id", "project_id", "thesis_id", "title", "title_source", "title_confirmed", "filename",
        "page_count", "size_mb", "status", "created_at", "updated_at",
        "analysis", "sections", "review", "preflight", "n_references", "active_page", "audit_decisions", "metadata", "ingestion"
    }
    out = {k: deepcopy(v) for k, v in doc.items() if k in keep}
    out["analysis_stale"] = _analysis_stale(doc)
    out["analysis_versions"] = _analysis_versions(doc)
    out["current_versions"] = {
        "software_version": SOFTWARE_VERSION,
        "algorithm_version": ALGORITHM_VERSION,
        "stable_engine_version": STABLE_ENGINE_VERSION,
        "extraction_revision": EXTRACTION_REVISION,
    }
    out["has_manual_corrections"] = bool(doc.get("corrections"))
    out["has_manual_work"] = _has_manual_work(doc)
    out["analysis_history_count"] = len(doc.get("analysis_history") or [])
    if include_refs:
        out["references"] = inherit_repeated_authors([_reference_with_parse(r) for r in doc.get("references", [])])
        out["corrections"] = deepcopy(doc.get("corrections", []))
    return out


def _discard_copy(path: Path) -> None:
    """Remove a project copy that was never registered, best effort.

    A file still held open by another process (or by a PDF reader that failed
    mid-open) would otherwise turn a rejected thesis into a failed request.
    """
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _create_document_from_pdf_path(source_path: Path, filename: str, ingestion: dict[str, Any] | None = None) -> dict[str, Any]:
    """Copy one validated PDF candidate into the project and register it.

    All import routes (browser upload, ZIP and Google Drive) converge here, so
    extraction behavior stays identical regardless of where the bytes came from.
    """
    doc_id = f"doc_{uuid.uuid4().hex[:12]}"
    filename = _safe_name(filename)
    path = UPLOAD_DIR / f"{doc_id}_{filename}"
    shutil.copyfile(source_path, path)
    try:
        preflight = pdf_preflight(path)
    except Exception as exc:
        unreadable = str(exc)
    else:
        unreadable = None
    if unreadable is not None:
        # Deliberately outside the except block: while the failed PyMuPDF
        # document is still reachable through the traceback, Windows keeps the
        # copy open and deleting it raises.  One broken thesis must not abort
        # the whole import.
        _discard_copy(path)
        raise HTTPException(400, f"No se pudo abrir {filename}: {unreadable}")
    detected_title, title_source = _detect_thesis_title(path, filename)
    try:
        return _register_prepared_document(path, filename, ingestion, preflight,
                                           detected_title, title_source, doc_id=doc_id)
    except DuplicateDocument as exc:
        _discard_copy(path)
        raise HTTPException(409, str(exc)) from exc


def _register_prepared_document(path, filename, ingestion, preflight,
                                detected_title, title_source, *, doc_id=None, extras=None):
    """Only the parent process owns document identities and project storage."""
    doc_id = doc_id or f"doc_{uuid.uuid4().hex[:12]}"
    doc = {
        "id": doc_id,
        "project_id": (ingestion or {}).get("project_id") or _upload_project(None)["id"],
        "title": detected_title,
        "title_source": title_source,
        "title_confirmed": False,
        "metadata": {},
        "filename": filename,
        "path": str(path),
        "page_count": preflight["page_count"],
        "size_mb": preflight["size_mb"],
        "preflight": preflight,
        "status": "uploaded",
        "created_at": _now(),
        "updated_at": _now(),
        "analysis": {},
        "sections": [],
        "references": [],
        "corrections": [],
        "n_references": 0,
        "active_page": 1,
        "ingestion": deepcopy(ingestion or {"source": "local_upload"}),
    }
    doc.update(extras or {})
    digest = doc["ingestion"].get("sha256") or sha256_file(path)
    doc["content_sha256"] = digest
    with STORE.lock:
        existing = next((d for d in STORE.docs().values()
                         if d.get("project_id") == doc["project_id"] and content_hash(d) == digest), None)
        if existing is not None:
            raise DuplicateDocument(existing, filename)
        doc["thesis_id"] = _next_thesis_id()
        STORE.docs()[doc_id] = doc
    return doc


def _ingest_zip_path(zip_path: Path, *, source_context: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    created: list[dict[str, Any]] = []
    warnings: list[str] = []
    with tempfile.TemporaryDirectory(prefix="biblioref_zip_", dir=str(WORKSPACE)) as tmpdir:
        try:
            pdfs, zip_warnings = extract_pdfs_from_zip(zip_path, tmpdir)
        except ZipIngestionError as exc:
            raise HTTPException(400, str(exc)) from exc
        warnings.extend(zip_warnings)
        for item in pdfs:
            provenance = item.provenance()
            provenance.update(source_context or {})
            display_archive = str((source_context or {}).get("archive_filename") or (source_context or {}).get("drive_file_name") or zip_path.name)
            chain = list(provenance.get("archive_chain") or [])
            if chain:
                chain[0] = display_archive
                provenance["archive_chain"] = chain
            # Preserve the fact that a Google Drive import came through a ZIP.
            if (source_context or {}).get("source") == "google_drive":
                provenance["source"] = "google_drive_zip"
            provenance["container_filename"] = display_archive
            try:
                created.append(_create_document_from_pdf_path(item.path, item.display_name, provenance))
            except HTTPException as exc:
                warnings.append(f"{item.display_name}: {exc.detail}")
    return created, warnings


def _ingest_saved_upload(path: Path, filename: str, *, source_context: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    low = filename.lower()
    if low.endswith(".pdf"):
        return [_create_document_from_pdf_path(path, filename, source_context or {"source": "local_upload"})], []
    if low.endswith(".zip"):
        ctx = {"source": "zip", "archive_filename": filename}
        ctx.update(source_context or {})
        return _ingest_zip_path(path, source_context=ctx)
    return [], [f"Formato omitido: {filename}"]


class Store:
    def __init__(self):
        self.lock = threading.RLock()
        self.data: dict[str, Any] = {"schema": 1, "documents": {}}
        self.load()

    def load(self):
        with self.lock:
            if STATE_FILE.exists():
                try:
                    self.data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
                except Exception:
                    self.data = {"schema": 1, "documents": {}}

    def save(self):
        with self.lock:
            tmp = STATE_FILE.with_suffix(".tmp")
            tmp.write_text(json.dumps(self.data, ensure_ascii=False, indent=2), encoding="utf-8")
            # Windows refuses to replace a file that anyone else has open, and
            # the project index is large enough that a backup tool, an antivirus
            # or a second window can be reading it exactly now.  Retry briefly
            # instead of losing the edit.
            for attempt in range(40):
                try:
                    tmp.replace(STATE_FILE)
                    return
                except PermissionError:
                    if attempt == 39:
                        raise
                    time.sleep(0.05)

    def docs(self):
        return self.data.setdefault("documents", {})

    def get(self, doc_id: str):
        doc = self.docs().get(doc_id)
        if not doc:
            raise HTTPException(404, "Tesis no encontrada")
        return doc

STORE = Store()
APP_TOKEN = uuid.uuid4().hex


def _ensure_projects() -> dict[str, Any]:
    """Migrate membership only. PDF bytes and all analytical fields stay intact."""
    with STORE.lock:
        projects = STORE.data.setdefault("projects", {})
        orphaned = [d for d in STORE.docs().values() if not d.get("project_id")]
        if orphaned:
            projects.setdefault(LEGACY_PROJECT_ID, {
                "id": LEGACY_PROJECT_ID, "name": "Proyecto anterior", "description": "",
                "created_at": min(d.get("created_at", _now()) for d in orphaned),
                "updated_at": max(d.get("updated_at", _now()) for d in orphaned),
                "group_loose_by_university": False,
            })
            for doc in orphaned:
                doc["project_id"] = LEGACY_PROJECT_ID
            STORE.data["schema"] = 2
            STORE.save()
        return projects


def _get_project(project_id: str) -> dict[str, Any]:
    project = _ensure_projects().get(project_id)
    if project is None:
        raise HTTPException(404, "Proyecto no encontrado")
    return project


def _upload_project(project_id: str | None) -> dict[str, Any]:
    if project_id:
        return _get_project(project_id)
    # Compatibility for previous clients; the new interface always sends an ID.
    with STORE.lock:
        return _ensure_projects().setdefault(LEGACY_PROJECT_ID, {
            "id": LEGACY_PROJECT_ID, "name": "Proyecto anterior", "description": "",
            "created_at": time.time(), "updated_at": time.time(),
            "group_loose_by_university": False,
        })


def _project_summary(project: dict[str, Any]) -> dict[str, Any]:
    docs = [d for d in STORE.docs().values() if d.get("project_id") == project["id"]]
    summary = deepcopy(project)
    summary["n_documents"] = len(docs)
    summary["n_references"] = sum(d.get("n_references", 0) for d in docs)
    summary["updated_at"] = max([project.get("updated_at", 0)] + [d.get("updated_at", 0) for d in docs])
    return summary


class ProjectPayload(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    description: str = Field(default="", max_length=1200)
    group_loose_by_university: bool = False


class ProjectUpdatePayload(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=160)
    description: str | None = Field(default=None, max_length=1200)
    group_loose_by_university: bool | None = None


def _refresh_reference_parse(ref: dict[str, Any]) -> dict[str, Any]:
    """Refresh downstream metadata without touching extraction provenance."""
    ref["parsed"] = parse_reference(_norm(ref.get("text") or ref.get("reference_text") or ""))
    return ref


def _reference_with_parse(ref: dict[str, Any]) -> dict[str, Any]:
    """Return parser truth plus an explicit, non-destructive human-review overlay."""
    row = deepcopy(ref)
    parsed = row.get("parsed") or {}
    if parsed.get("parser_version") != PARSER_VERSION:
        row["parsed"] = parse_reference(_norm(row.get("text") or row.get("reference_text") or ""))
    parsed = row.get("parsed") or {}
    review = deepcopy(row.get("parsed_review") or {})
    overrides = dict(review.get("overrides") or {})
    overrides = {key: value for key, value in overrides.items() if key in PARSED_REVIEW_FIELDS}
    effective = deepcopy(parsed)
    for key, value in overrides.items():
        effective[key] = deepcopy(value)
    review["overrides"] = overrides
    review["status"] = review.get("status") or ("edited" if overrides else "parser_original")
    row["parsed_review"] = review
    row["effective_parsed"] = effective
    return row


RESOLUTION_SIDECAR_KEYS = (
    "resolver_query", "resolver_batch", "resolver_candidates", "consensus",
    "semantic", "rag_result", "rag_adjudication", "adjudication",
    "resolution_decision", "resolution_run",
)


def _reset_reference_resolution(ref: dict[str, Any], *, preserve_human: bool = True) -> None:
    """Invalidate downstream evidence whenever its bibliographic input changes."""
    for key in RESOLUTION_SIDECAR_KEYS:
        ref.pop(key, None)
    ref.pop("llm_chat", None)
    if not preserve_human:
        ref.pop("human_resolution", None)


PARSED_EXPORT_FIELDS = [
    "parsed_parser_version", "parsed_parser_status", "parsed_parse_style", "parsed_author", "parsed_authors",
    "parsed_year", "parsed_publication_type", "parsed_title", "parsed_journal",
    "parsed_publisher", "parsed_publisher_observed", "parsed_source_publisher", "parsed_doi",
]

def _flat_parse_export(ref: dict[str, Any]) -> dict[str, Any]:
    return flatten_parsed((_reference_with_parse(ref).get("parsed") or {}))


def _record_to_ui(record: dict[str, Any], index: int, source: str = "automatic") -> dict[str, Any]:
    text = _norm(record.get("reference_text") or record.get("text") or "")
    row = {
        "id": f"ref_{uuid.uuid4().hex[:12]}",
        "order": index,
        "text": text,
        "original_text": text,
        "start_page": record.get("reference_start_page") or record.get("zone_start_page"),
        "end_page": record.get("reference_end_page") or record.get("zone_end_page") or record.get("reference_start_page"),
        "start_x": record.get("reference_start_x"),
        "zone_id": record.get("zone_id"),
        "style": record.get("style"),
        "boundary_evidence": record.get("boundary_evidence", ""),
        "audit_flags": record.get("audit_flags", ""),
        "status": "extracted",
        "source": source,
        "edited": False,
        "deleted": False,
    }
    return _refresh_reference_parse(row)


def _zones_to_sections(zones: list[dict[str, Any]], source: str = "automatic") -> list[dict[str, Any]]:
    sections = []
    for i, zone in enumerate(zones, 1):
        sections.append({
            "id": f"sec_{uuid.uuid4().hex[:10]}",
            "zone_id": zone.get("zone_id", i),
            "label": zone.get("heading") or zone.get("zone_heading") or f"Sección de referencias {i}",
            "start_page": zone.get("start_page") or zone.get("zone_start_page"),
            "start_y": zone.get("start_y"),
            "end_page": zone.get("end_page") or zone.get("zone_end_page"),
            "end_y": zone.get("end_y"),
            "role": zone.get("zone_role", ""),
            "source": source,
            "confirmed": source == "manual",
        })
    return sections


def _analysis_summary(result: dict[str, Any]) -> dict[str, Any]:
    audit = result.get("post_extraction_audit") or {}
    plaus = audit.get("output_plausibility") or {}
    legacy_coverage = audit.get("document_coverage") or {}
    coverage = result.get("coverage_report") or legacy_coverage
    detection = result.get("detection_decision") or {}
    shadow = result.get("adaptive_shadow") or {}
    status = audit.get("status") or ("review" if detection.get("requires_confirmation") else "ok")
    reasons = list(audit.get("review_reasons") or [])
    if shadow.get("overall_status") == "manual_review_required" and "adaptive_review" not in reasons:
        reasons.append("adaptive_review")
    # More paragraphs left out than references taken (T0145: 19 taken, 72 left out) is an incomplete
    # extraction even when the engine's own audit found nothing wrong.
    n_missing = int((coverage.get("summary") or {}).get("n_possible_missing") or 0)
    if n_missing >= 10 and n_missing > len(result.get("records") or []) and "coverage_gap" not in reasons:
        reasons.append("coverage_gap")
    human = []
    reason_map = {
        "zero_records": "No se extrajo ninguna referencia aunque el documento muestra señales bibliográficas.",
        "adaptive_review": "El análisis adicional encontró una estructura que conviene comprobar visualmente.",
        "implausible_record_count": "La cantidad de referencias parece inusual para la extensión del documento.",
        "low_reference_density": "La cobertura bibliográfica puede ser incompleta.",
        "coverage_gap": "Quedaron afuera más párrafos de la bibliografía que referencias extraídas: la extracción probablemente está incompleta.",
    }
    for r in reasons:
        human.append(reason_map.get(r, str(r).replace("_", " ")))
    if legacy_coverage.get("possible_incomplete_coverage"):
        human.append("Puede haber secciones de referencias no cubiertas por la extracción efectiva.")
    if (coverage.get("summary") or {}).get("n_possible_missing", 0):
        human.append(f"El control de cobertura encontró {(coverage.get('summary') or {}).get('n_possible_missing')} posible(s) referencia(s) no representada(s).")
    if (coverage.get("summary") or {}).get("n_page_boundaries_review", 0):
        human.append(f"Hay {(coverage.get('summary') or {}).get('n_page_boundaries_review')} borde(s) de página para revisar.")
    return {
        "status": status,
        "needs_review": bool(status not in ("ok", "looks_consistent") or detection.get("requires_confirmation") or reasons),
        "messages": human,
        "detection_method": detection.get("method"),
        "requires_confirmation": bool(detection.get("requires_confirmation")),
        "plausibility": plaus,
        "coverage": coverage,
        "legacy_coverage": legacy_coverage,
        "recommendations": audit.get("recommendations") or [],
        "rescue_menu": audit.get("menu") or [],
        "software_version": result.get("software_version", SOFTWARE_VERSION),
        "algorithm_version": result.get("algorithm_version", ALGORITHM_VERSION),
        "stable_engine_version": result.get("stable_engine_version", STABLE_ENGINE_VERSION),
        "extraction_revision": EXTRACTION_REVISION,
    }


def _renumber(doc: dict[str, Any]):
    visible = [r for r in doc.get("references", []) if not r.get("deleted")]
    for i, ref in enumerate(visible, 1):
        ref["order"] = i
    doc["n_references"] = len(visible)



def _log_correction(doc: dict[str, Any], kind: str, payload: dict[str, Any]):
    doc.setdefault("corrections", []).append({
        "id": f"corr_{uuid.uuid4().hex[:12]}",
        "kind": kind,
        "timestamp": _now(),
        **deepcopy(payload),
    })


def _analyze_document(doc: dict[str, Any], rescue_profile: str = "stable", *, prepared_result=None):
    path = Path(doc["path"])
    if doc.get("references") or doc.get("sections"):
        _archive_analysis_snapshot(doc, reason=f"reprocess_to_{SOFTWARE_VERSION}")
    result = prepared_result if prepared_result is not None else process_pdf(
        path,
        scope="all_reference_sections",
        adaptive_mode="shadow",
        rescue_profile=rescue_profile,
        rescue_audit=True,
    )
    result_path = RESULT_DIR / f"{doc['id']}_result.json"
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    doc["analysis"] = _analysis_summary(result)
    doc["sections"] = _zones_to_sections(list(result.get("zones") or []), "automatic")
    doc["references"] = [
        _record_to_ui(r, i, "automatic") for i, r in enumerate(result.get("records") or [], 1)
    ]
    doc["n_references"] = len(doc["references"])
    doc["status"] = "review" if doc["analysis"]["needs_review"] else "ready"
    doc["updated_at"] = _now()
    doc["result_file"] = str(result_path)
    _renumber(doc)


def _manual_sections(doc: dict[str, Any], specs: list[dict[str, Any]]):
    all_records: list[dict[str, Any]] = []
    sections: list[dict[str, Any]] = []
    for zone_id, raw in enumerate(specs, 1):
        spec = ManualZoneSpec(
            start_page=int(raw["start_page"]),
            start_y=float(raw.get("start_y") or 0),
            end_page=int(raw["end_page"]),
            end_y=(float(raw["end_y"]) if raw.get("end_y") not in (None, "") else None),
            heading_text=str(raw.get("label") or f"Sección manual {zone_id}"),
            note=str(raw.get("note") or ""),
        )
        processed = process_manual_zone(doc["path"], spec, zone_id=zone_id)
        sections.extend(_zones_to_sections([processed["zone"]], "manual"))
        all_records.extend(processed["records"])
    before = len([r for r in doc.get("references", []) if not r.get("deleted")])
    doc["sections"] = sections
    doc["references"] = [_record_to_ui(r, i, "manual") for i, r in enumerate(all_records, 1)]
    _renumber(doc)
    doc["status"] = "review"
    doc["updated_at"] = _now()
    try:
        cov_result = {
            "zones": [{
                "zone_id": s.get("zone_id", i + 1),
                "start_page": s.get("start_page"), "start_y": s.get("start_y"),
                "end_page": s.get("end_page"), "end_y": s.get("end_y"),
            } for i, s in enumerate(doc.get("sections", []))],
            "records": [{
                "zone_id": r.get("zone_id"), "ref_in_zone": r.get("order"),
                "reference_text": r.get("text"), "reference_start_page": r.get("start_page"),
                "reference_end_page": r.get("end_page"),
            } for r in doc.get("references", []) if not r.get("deleted")],
        }
        doc.setdefault("analysis", {})["coverage"] = build_coverage_report(doc["path"], cov_result)
    except Exception as exc:
        doc.setdefault("analysis", {})["coverage"] = {"status":"unavailable","error":str(exc),"possible_missing":[],"page_boundaries":[],"merged_suspects":[],"summary":{}}
    _log_correction(doc, "manual_sections_reprocessed", {
        "before_reference_count": before,
        "after_reference_count": doc["n_references"],
        "sections": specs,
    })


class AnalyzePayload(BaseModel):
    rescue_profile: str = "stable"
    worker: bool = False


class MetadataPayload(BaseModel):
    title: str | None = Field(default=None, max_length=500)
    author: str | None = Field(default=None, max_length=300)
    year: str | None = Field(default=None, max_length=40)
    university: str | None = Field(default=None, max_length=300)
    faculty: str | None = Field(default=None, max_length=300)
    program: str | None = Field(default=None, max_length=300)
    repository: str | None = Field(default=None, max_length=300)
    repository_url: str | None = Field(default=None, max_length=1000)
    area: str | None = Field(default=None, max_length=300)
    country: str | None = Field(default=None, max_length=120)


class BulkMetadataPayload(BaseModel):
    document_ids: list[str] = Field(min_length=1)
    fields: dict[str, str] = Field(default_factory=dict)


class RefEditPayload(BaseModel):
    text: str
    status: str | None = None


class ParsedReviewPayload(BaseModel):
    overrides: dict[str, Any] = Field(default_factory=dict)
    unset_fields: list[str] = Field(default_factory=list)


class ResolutionRunPayload(BaseModel):
    layers: list[str] = Field(default_factory=lambda: ["crossref", "openalex", "consensus", "decision"], max_length=4)


class HumanResolutionPayload(BaseModel):
    decision: str = "candidate"
    candidate_index: int | None = None
    note: str = Field(default="", max_length=1000)


class SplitPayload(BaseModel):
    parts: list[str] = Field(min_length=2)


class MergePayload(BaseModel):
    with_ref_id: str


class InsertPayload(BaseModel):
    text: str = ""
    start_page: int | None = None
    end_page: int | None = None
    start_y: float | None = None
    end_y: float | None = None
    after_ref_id: str | None = None
    before_ref_id: str | None = None
    source: str = "human_insert"


class ReferenceCandidatePayload(BaseModel):
    page: int
    y: float
    x: float | None = None
    end_page: int | None = None
    end_y: float | None = None
    section_label: str = "Referencias"
    margin_pt: float = Field(default=PDF_CLICK_MARGIN_PT, ge=0.0, le=72.0)


class ManualSectionsPayload(BaseModel):
    sections: list[dict[str, Any]] = Field(min_length=1)


class AuditDecisionPayload(BaseModel):
    decision: str
    note: str = ""


class DriveImportPayload(BaseModel):
    file_ids: list[str] = Field(min_length=1, max_length=200)
    recursive_folders: bool = True
    project_id: str | None = None


app = FastAPI(title="BiblioRef", version=str(SOFTWARE_VERSION))

@app.middleware("http")
async def local_write_guard(request, call_next):
    if request.url.path.startswith("/api/") and request.method.upper() in {"POST", "PATCH", "PUT", "DELETE"}:
        if request.headers.get("X-BiblioRef-Token") != APP_TOKEN:
            return Response(content=json.dumps({"detail": "Solicitud local no autorizada"}), status_code=403, media_type="application/json")
    response = await call_next(request)
    # This is a local desktop-style app. Stale browser assets are more harmful
    # than the tiny performance gain from caching, especially while the UI and
    # extraction revision evolve together.
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


@app.get("/api/projects")
def list_projects():
    projects = [_project_summary(p) for p in _ensure_projects().values()]
    projects.sort(key=lambda p: max(p.get("updated_at", 0), p.get("last_opened_at", 0)), reverse=True)
    return {"projects": projects, "version": PROJECTS_VERSION}


@app.post("/api/projects")
def create_project(payload: ProjectPayload):
    name = _norm(payload.name)
    if not name:
        raise HTTPException(422, "Escribí un nombre para el proyecto")
    with STORE.lock:
        projects = _ensure_projects()
        project = {"id": f"project_{uuid.uuid4().hex[:12]}", "name": name,
                   "description": payload.description.strip(), "created_at": time.time(),
                   "updated_at": time.time(), "group_loose_by_university": payload.group_loose_by_university}
        projects[project["id"]] = project
        STORE.data["schema"] = 2
        STORE.save()
    return _project_summary(project)


@app.patch("/api/projects/{project_id}")
def update_project(project_id: str, payload: ProjectUpdatePayload):
    with STORE.lock:
        project = _get_project(project_id)
        if payload.name is not None:
            name = _norm(payload.name)
            if not name:
                raise HTTPException(422, "Escribí un nombre para el proyecto")
            project["name"] = name
        if payload.description is not None:
            project["description"] = payload.description.strip()
        if payload.group_loose_by_university is not None:
            project["group_loose_by_university"] = payload.group_loose_by_university
        project["updated_at"] = time.time()
        STORE.save()
    return _project_summary(project)


@app.post("/api/projects/{project_id}/open")
def open_project(project_id: str):
    with STORE.lock:
        project = _get_project(project_id)
        project["last_opened_at"] = time.time()
        STORE.save()
    return _project_summary(project)


@app.get("/api/projects/{project_id}/tree")
def get_project_tree(project_id: str):
    project = _get_project(project_id)
    docs = [d for d in STORE.docs().values() if d.get("project_id") == project_id]
    return {"project_id": project_id, "nodes": project_tree(project, docs)}


@app.get("/", response_class=HTMLResponse)
def index():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.get("/api/health")
def health():
    return {
        "ok": True,
        "software_version": SOFTWARE_VERSION,
        "algorithm_version": ALGORITHM_VERSION,
        "stable_engine_version": STABLE_ENGINE_VERSION,
        "workspace": str(WORKSPACE),
        "session_token": APP_TOKEN,
        "parser_version": PARSER_VERSION,
        "parser_status": PARSER_STATUS,
        "app_version": APP_VERSION,
        "zip_ingestion_version": ZIP_INGESTION_VERSION,
    }


def _processing_layer_manifest() -> list[dict[str, Any]]:
    """Describe pipeline capabilities truthfully without implying they already ran."""
    return [
        {"id": "extraction", "label": "Extracción", "mode": "core", "available": True, "default_enabled": True, "interactive_execution": True, "version": EXTRACTION_REVISION, "description": "Detecta y segmenta la bibliografía desde el PDF."},
        {"id": "parser", "label": "Parser", "mode": "core", "available": True, "default_enabled": True, "interactive_execution": True, "version": PARSER_VERSION, "description": "Separa autor, año, tipo, título, revista/editorial y DOI."},
        {"id": "styleprofiler", "label": "StyleProfiler", "mode": "shadow", "available": True, "default_enabled": False, "interactive_execution": False, "version": STYLE_PROFILER_VERSION, "description": "Implementado en shadow; todavía no se ejecuta desde esta pantalla."},
        {"id": "crossref", "label": "Crossref", "mode": "external", "available": True, "default_enabled": True, "interactive_execution": True, "version": "metadata-backend-v0.1", "description": "Busca candidatos en Crossref desde la referencia seleccionada."},
        {"id": "openalex", "label": "OpenAlex", "mode": "external", "available": True, "default_enabled": True, "interactive_execution": True, "version": "metadata-backend-v0.1", "description": "Busca candidatos en OpenAlex desde la referencia seleccionada."},
        {"id": "consensus", "label": "Consensus", "mode": "resolver", "available": True, "default_enabled": True, "interactive_execution": True, "version": "metadata-consensus-v0.2", "description": "Compara evidencia determinística entre candidatos y fuentes."},
        {"id": "semantic", "label": "Semantic", "mode": "deferred_local", "available": False, "default_enabled": False, "interactive_execution": False, "version": "semantic-resolver-v0.1", "description": "Implementado experimentalmente, pero desactivado en esta versión: no se ejecutan modelos locales."},
        {"id": "rag", "label": "RAG", "mode": "deferred_local", "available": False, "default_enabled": False, "interactive_execution": False, "version": "bibliographic-rag-v0.1", "description": "Implementado experimentalmente, pero desactivado en esta versión junto con el runtime local."},
        {"id": "llm", "label": "LLM adjudicador", "mode": "deferred_local", "available": False, "default_enabled": False, "interactive_execution": False, "version": "local-ai-adjudicator-v0.2-rag", "description": "Implementado experimentalmente, pero desactivado: la app actual no ejecuta LLM locales."},
        {"id": "decision", "label": "ResolutionDecision", "mode": "resolver", "available": True, "default_enabled": True, "interactive_execution": True, "version": "resolution-decision-v0.1", "description": "Integra evidencia determinística de consenso sin escribir sobre el parser."},
        {"id": "human_review", "label": "Revisión humana", "mode": "review", "available": True, "default_enabled": True, "interactive_execution": True, "version": UI_WORKBENCH_VERSION, "description": "Corrige campos y confirma candidatos en una capa humana explícita."},
    ]


@app.get("/api/capabilities")
def capabilities():
    from biblioref.pdf_workers import worker_count
    return {
        "processing": {"max_workers": worker_count(), "mode": "processes", "parallel_unit": "document"},
        "imports": {"pdf": True, "multiple_pdf": True, "folder_desktop": True, "zip": True, "nested_zip": True, "docx": False, "ocr": False},
        "exports": {"xlsx": Workbook is not None, "csv": True, "json": True, "ris": False, "bibtex": False, "csl_json": False},
        "integrations": {"google_drive": {"mode": "read_only", "available": gdrive.GOOGLE_DRIVE_LIBS_AVAILABLE}, "zotero_local": "planned"},
        "calibration": {"manual_sections": True, "vertical_boundaries": True, "edit": True, "split": True, "merge": True, "insert": True, "delete": True, "source_trace": True, "parsed_field_review": True},
        "parsing": {"enabled": True, "parser_version": PARSER_VERSION, "status": PARSER_STATUS, "fields": ["author", "authors", "year", "publication_type", "title", "journal", "publisher", "publisher_observed", "source_publisher", "doi"], "reviewable_fields": list(PARSED_REVIEW_FIELDS)},
        "layers": _processing_layer_manifest(),
        "local_runtime": {"enabled": False, "policy": "deferred", "active_components": []},
        "ui_workbench_version": UI_WORKBENCH_VERSION,
    }


@app.post("/api/documents")
async def upload_documents(files: list[UploadFile] = File(...), project_id: str | None = Form(None), stream_progress: bool = Form(False)):
    project = _upload_project(project_id)
    if stream_progress:
        import sys
        from biblioref.import_stream import stage_and_stream
        return await stage_and_stream(sys.modules[__name__], files, project)
    created: list[dict[str, Any]] = []
    warnings: list[str] = []
    for uploaded in files:
        if not uploaded.filename:
            continue
        try:
            logical_path = relative_path(uploaded.filename)
        except ValueError as exc:
            warnings.append(f"Archivo omitido: {exc}")
            continue
        filename = _safe_name(logical_path.rsplit("/", 1)[-1])
        if not filename.lower().endswith((".pdf", ".zip")):
            warnings.append(f"Formato omitido: {filename}")
            continue
        suffix = ".zip" if filename.lower().endswith(".zip") else ".pdf"
        with tempfile.NamedTemporaryFile(prefix="biblioref_upload_", suffix=suffix, dir=str(WORKSPACE), delete=False) as out:
            tmp_path = Path(out.name)
            while True:
                chunk = await uploaded.read(4 * 1024 * 1024)
                if not chunk:
                    break
                out.write(chunk)
        try:
            docs, local_warnings = _ingest_saved_upload(tmp_path, filename, source_context={
                "source": "folder" if "/" in logical_path else ("zip" if suffix == ".zip" else "local_upload"),
                "relative_path": logical_path, "project_id": project["id"],
            })
            created.extend(docs)
            warnings.extend(local_warnings)
        except HTTPException as exc:
            warnings.append(f"{logical_path}: {exc.detail}")
        finally:
            tmp_path.unlink(missing_ok=True)
    if not created:
        detail = "No se encontró ningún PDF para cargar"
        if warnings:
            detail += ": " + "; ".join(warnings[:5])
        raise HTTPException(400, detail)
    project["updated_at"] = time.time()
    STORE.save()
    return {"documents": [_public_doc(d, True) for d in created], "warnings": warnings}



@app.get("/api/integrations/google-drive/status")
def google_drive_status():
    return gdrive.status()


@app.post("/api/integrations/google-drive/client-secret")
async def google_drive_client_secret(file: UploadFile = File(...)):
    raw = await file.read(2 * 1024 * 1024)
    if len(raw) >= 2 * 1024 * 1024:
        raise HTTPException(400, "El archivo OAuth es inesperadamente grande")
    try:
        path = gdrive.save_client_secret_bytes(raw)
    except gdrive.GoogleDriveError as exc:
        raise HTTPException(400, str(exc)) from exc
    return {"ok": True, "configured": True, "filename": path.name}


@app.post("/api/integrations/google-drive/connect")
def google_drive_connect():
    try:
        gdrive.connect()
        return gdrive.status()
    except gdrive.GoogleDriveError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, f"No se pudo completar la conexión con Google Drive: {exc}") from exc


@app.post("/api/integrations/google-drive/disconnect")
def google_drive_disconnect():
    gdrive.disconnect()
    return {"ok": True, "connected": False}


@app.get("/api/integrations/google-drive/list")
def google_drive_list(folder_id: str = "root"):
    try:
        return gdrive.list_folder(folder_id or "root")
    except gdrive.GoogleDriveError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, f"No se pudo leer Google Drive: {exc}") from exc


@app.get("/api/integrations/google-drive/search")
def google_drive_search(q: str, limit: int = 100):
    try:
        return gdrive.search_items(q, max_results=limit)
    except gdrive.GoogleDriveError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, f"No se pudo buscar en Google Drive: {exc}") from exc


@app.post("/api/integrations/google-drive/import")
def google_drive_import(payload: DriveImportPayload):
    project = _upload_project(payload.project_id)
    created: list[dict[str, Any]] = []
    warnings: list[str] = []
    seen: set[str] = set()
    try:
        with tempfile.TemporaryDirectory(prefix="biblioref_drive_", dir=str(WORKSPACE)) as tmpdir:
            tmpdir_path = Path(tmpdir)
            queue: list[dict[str, Any]] = []
            for requested_id in payload.file_ids:
                raw = gdrive.get_item(requested_id)
                effective_mime = gdrive.effective_mime(raw)
                effective_id = gdrive.effective_id(raw)
                if effective_mime == gdrive.FOLDER_MIME:
                    if not payload.recursive_folders:
                        continue
                    queue.extend(list(gdrive.iter_folder_files(effective_id, recursive=True)))
                else:
                    queue.append({
                        "id": requested_id,
                        "effective_id": effective_id,
                        "name": raw.get("name") or requested_id,
                        "mime_type": effective_mime,
                        "is_folder": False,
                        "is_shortcut": raw.get("mimeType") == gdrive.SHORTCUT_MIME,
                    })
            for item in queue:
                file_id = str(item.get("effective_id") or item.get("id") or "")
                if not file_id or file_id in seen:
                    continue
                seen.add(file_id)
                name = _safe_name(str(item.get("name") or file_id))
                low = name.lower()
                if not low.endswith((".pdf", ".zip")) and item.get("mime_type") != gdrive.PDF_MIME:
                    warnings.append(f"Drive omitió formato no compatible: {name}")
                    continue
                if item.get("mime_type") == gdrive.PDF_MIME and not low.endswith(".pdf"):
                    name += ".pdf"
                    low = name.lower()
                suffix = ".zip" if low.endswith(".zip") else ".pdf"
                local = tmpdir_path / f"{uuid.uuid4().hex}{suffix}"
                gdrive.download_blob(file_id, local)
                context = {
                    "source": "google_drive",
                    "project_id": project["id"],
                    "drive_file_id": file_id,
                    "drive_file_name": name,
                    "google_drive_integration_version": gdrive.GOOGLE_DRIVE_INTEGRATION_VERSION,
                }
                docs, local_warnings = _ingest_saved_upload(local, name, source_context=context)
                created.extend(docs)
                warnings.extend(local_warnings)
    except gdrive.GoogleDriveError as exc:
        raise HTTPException(400, str(exc)) from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(502, f"No se pudo importar desde Google Drive: {exc}") from exc
    if not created:
        raise HTTPException(400, "La selección de Google Drive no contenía PDF importables")
    STORE.save()
    return {"documents": [_public_doc(d, True) for d in created], "warnings": warnings}


@app.get("/api/documents")
def list_documents(project_id: str | None = None):
    _ensure_projects()
    if project_id is not None:
        _get_project(project_id)
    docs = sorted(STORE.docs().values(), key=lambda d: d.get("created_at", 0))
    if project_id is not None:
        docs = [d for d in docs if d.get("project_id") == project_id]
    changed = False
    for d in docs:
        changed = _ensure_doc_identity(d) or changed
    if changed:
        STORE.save()
    return {"documents": [_public_doc(d, False) for d in docs]}


@app.get("/api/documents/{doc_id}")
def get_document(doc_id: str):
    doc = STORE.get(doc_id)
    if _ensure_doc_identity(doc):
        STORE.save()
    return _public_doc(doc, True)


@app.patch("/api/documents/{doc_id}/metadata")
def update_document_metadata(doc_id: str, payload: MetadataPayload):
    doc = STORE.get(doc_id)
    _ensure_doc_identity(doc)
    data = payload.model_dump(exclude_unset=True) if hasattr(payload, "model_dump") else payload.dict(exclude_unset=True)
    changes: dict[str, Any] = {}
    if "title" in data:
        title = _norm(data.get("title") or "")
        if not title:
            raise HTTPException(422, "El título no puede quedar vacío")
        before = doc.get("title")
        if title != before:
            doc["title"] = title
            doc["title_source"] = "human"
            doc["title_confirmed"] = True
            changes["title"] = {"before": before, "after": title}
    meta = doc.setdefault("metadata", {})
    for key in ("author", "year", "university", "faculty", "program", "repository", "repository_url", "area", "country"):
        if key not in data:
            continue
        value = _norm(str(data.get(key) or ""))
        before = meta.get(key)
        if value:
            meta[key] = value
        else:
            meta.pop(key, None)
        if value != (before or ""):
            changes[key] = {"before": before, "after": value or None}
    if changes:
        doc["updated_at"] = _now()
        _log_correction(doc, "edit_thesis_metadata", {"changes": changes})
    STORE.save()
    return _public_doc(doc, True)


@app.patch("/api/documents/bulk-metadata")
def update_documents_bulk_metadata(payload: BulkMetadataPayload):
    allowed = {"university", "faculty", "program", "repository", "repository_url", "area", "country"}
    fields = {k: _norm(str(v or "")) for k, v in (payload.fields or {}).items() if k in allowed and _norm(str(v or ""))}
    if not fields:
        raise HTTPException(422, "No hay datos comunes para aplicar")
    updated = []
    for doc_id in dict.fromkeys(payload.document_ids):
        doc = STORE.get(doc_id)
        _ensure_doc_identity(doc)
        meta = doc.setdefault("metadata", {})
        changes = {}
        for key, value in fields.items():
            before = meta.get(key)
            if before != value:
                meta[key] = value
                changes[key] = {"before": before, "after": value}
        if changes:
            doc["updated_at"] = _now()
            _log_correction(doc, "bulk_edit_thesis_metadata", {"changes": changes})
        updated.append(_public_doc(doc, False))
    STORE.save()
    return {"documents": updated, "updated_count": len(updated)}


@app.delete("/api/documents/{doc_id}")
def delete_document(doc_id: str):
    doc = STORE.get(doc_id)
    Path(doc.get("path", "")).unlink(missing_ok=True)
    if doc.get("original_path"):  # the PDF as imported, before OCR
        Path(doc["original_path"]).unlink(missing_ok=True)
    if doc.get("result_file"):
        Path(doc["result_file"]).unlink(missing_ok=True)
    STORE.docs().pop(doc_id, None)
    STORE.save()
    return {"ok": True}


ACTIVE_ANALYSES: set[str] = set()


@app.post("/api/documents/{doc_id}/analyze")
async def analyze_document(doc_id: str, payload: AnalyzePayload):
    doc = STORE.get(doc_id)
    if doc_id in ACTIVE_ANALYSES:
        raise HTTPException(409, "Esta tesis ya se está procesando.")
    ACTIVE_ANALYSES.add(doc_id)
    try:
        doc["status"] = "processing"
        STORE.save()
        if payload.worker:
            from concurrent.futures.process import BrokenProcessPool
            from biblioref.pdf_workers import get_pool, analyze_pdf, discard_broken_pool
            pool = get_pool()
            try:
                result = await asyncio.wrap_future(pool.submit(analyze_pdf, doc["path"], payload.rescue_profile))
            except BrokenProcessPool:
                discard_broken_pool(pool)
                raise RuntimeError("El proceso de PDF se interrumpió. Podés volver a intentar esta tesis.")
            # Persist the authoritative core result in the parent process.
            _analyze_document(doc, payload.rescue_profile, prepared_result=result)
        else:
            _analyze_document(doc, payload.rescue_profile)
    except Exception as exc:
        doc["status"] = "error"
        doc["analysis"] = {"needs_review": True, "messages": [str(exc)]}
        STORE.save()
        raise HTTPException(500, f"Error al analizar la tesis: {exc}")
    finally:
        ACTIVE_ANALYSES.discard(doc_id)
    STORE.save()
    return _public_doc(doc, True)


@app.post("/api/documents/{doc_id}/manual-sections")
def apply_manual_sections(doc_id: str, payload: ManualSectionsPayload):
    doc = STORE.get(doc_id)
    try:
        _manual_sections(doc, payload.sections)
    except Exception as exc:
        raise HTTPException(400, str(exc))
    STORE.save()
    return _public_doc(doc, True)


@app.get("/api/documents/{doc_id}/page/{page_number}.png")
def page_png(doc_id: str, page_number: int, zoom: float = 1.35):
    doc = STORE.get(doc_id)
    try:
        png = render_page_png(doc["path"], page_number, zoom=min(max(zoom, 0.2), 2.5))
    except Exception as exc:
        raise HTTPException(400, str(exc))
    return Response(content=png, media_type="image/png", headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"})


@app.get("/api/documents/{doc_id}/page/{page_number}/lines")
def page_lines(doc_id: str, page_number: int):
    doc = STORE.get(doc_id)
    try:
        lines = extract_layout_lines_range(doc["path"], page_number, page_number)
        width, height = page_dimensions(doc["path"], page_number)
    except Exception as exc:
        raise HTTPException(400, str(exc))
    return {
        "page": page_number, "width": width, "height": height,
        "lines": [{"text": l.text, "x0": l.x0, "y0": l.y0, "x1": l.x1, "y1": l.y1} for l in lines]
    }


@app.get("/api/documents/{doc_id}/section-preview")
def section_preview(
    doc_id: str,
    start_page: int,
    end_page: int,
    start_y: float = 0.0,
    end_y: float | None = None,
):
    """Return page geometry for a continuous, cropped bibliography preview."""
    doc = STORE.get(doc_id)
    page_count = int(doc.get("page_count") or 0)
    if start_page < 1 or end_page < start_page or end_page > page_count:
        raise HTTPException(400, "Límites de sección inválidos")
    if end_page - start_page > 180:
        raise HTTPException(400, "La vista continua admite hasta 181 páginas por sección")
    pages = []
    try:
        for page in range(start_page, end_page + 1):
            width, height = page_dimensions(doc["path"], page)
            crop_top = max(0.0, float(start_y)) if page == start_page else 0.0
            crop_bottom = float(end_y) if page == end_page and end_y is not None else float(height)
            crop_bottom = min(float(height), max(crop_top + 1.0, crop_bottom))
            pages.append({
                "page": page, "width": float(width), "height": float(height),
                "crop_top": crop_top, "crop_bottom": crop_bottom,
            })
    except Exception as exc:
        raise HTTPException(400, str(exc))
    return {"start_page": start_page, "end_page": end_page, "pages": pages}


def _reference_match_score(a: str, b: str) -> float:
    aa = re.sub(r"[^a-z0-9áéíóúüñ]+", " ", _norm(a).lower()).strip()[:260]
    bb = re.sub(r"[^a-z0-9áéíóúüñ]+", " ", _norm(b).lower()).strip()[:260]
    if not aa or not bb:
        return 0.0
    if aa == bb:
        return 1.0
    return SequenceMatcher(None, aa, bb).ratio()


def _candidate_key(text: str) -> str:
    return re.sub(r"[^a-z0-9áéíóúüñ]+", " ", _norm(text).lower()).strip()


def _line_opens_record(line_text: str, record_text: str) -> bool:
    """Conservative prefix match used only to locate a segmenter's first boundary."""
    line = _candidate_key(line_text)
    record = _candidate_key(record_text)
    if not line or not record:
        return False
    # OCR often changes punctuation/spaces, so compare a lexical prefix rather
    # than raw text. Six tokens are enough to distinguish adjacent references.
    tokens = line.split()
    prefix = " ".join(tokens[: min(6, len(tokens))])
    if len(prefix) >= 8 and record.startswith(prefix):
        return True
    record_prefix = " ".join(record.split()[: min(6, len(record.split()))])
    return bool(record_prefix and line.startswith(record_prefix) and len(record_prefix) >= 8)


def _compact_ocr_word(text: str) -> str:
    return re.sub(r"[^a-záéíóúüñ]+", "", _norm(text).lower())


def _looks_like_bibliography_heading(text: str) -> bool:
    compact = _compact_ocr_word(text)
    return compact in {
        "bibliografia", "bibliografía", "bibliography", "references", "reference", "referencias",
        "referencies", "referenciasbibliograficas", "listofreferences", "referencelist",
    }


def _author_opening_score(lines, idx: int) -> float:
    """Presentation/manual-click heuristic, not an extraction rule.

    It only decides which physical line the user most likely meant when clicking
    near a reference start. A deferred year on the next line (RIEA/Alonso) is
    intentionally positive evidence here.
    """
    text = _norm(lines[idx].text)
    if not text or _looks_like_bibliography_heading(text):
        return -8.0
    low = text.lower()
    score = 0.0
    if re.search(r"\((?:18|19|20)\d{2}[a-z]?\)", text, re.I):
        score += 5.0
    if re.search(r"(?:^|[,;]\s*)(?:18|19|20)\d{2}[a-z]?(?:[,.:)]|\s)", text, re.I):
        score += 3.5
    if re.search(r"\b(?:eds?\.?|coords?\.?|comps?\.?)\b", text, re.I):
        score += 2.0
    if ";" in text:
        score += 2.2
    if "," in text:
        score += 1.2
    if re.search(r"\b[A-ZÁÉÍÓÚÜÑ][a-záéíóúüñ]{2,}\b", text):
        score += 0.8
    if re.search(r"\b[A-ZÁÉÍÓÚÜÑ]\.?\s*(?:[,;]|$)", text):
        score += 1.0
    # Deferred-year author openings: look ahead without swallowing the next ref.
    lookahead = " ".join(_norm(lines[j].text) for j in range(idx + 1, min(len(lines), idx + 3)))
    if re.search(r"\((?:18|19|20)\d{2}[a-z]?\)", lookahead, re.I) and ("," in text or ";" in text):
        score += 4.2
    # Strong continuation cues should not win a manual start snap.
    if re.match(r"^[a-záéíóúüñ]", text):
        score -= 3.0
    if re.match(r"^[)\].,:;]", text):
        score -= 2.5
    return score


def _select_click_start_line(lines, page: int, y: float, x: float | None, margin_pt: float):
    """Snap a manual click to a plausible bibliographic opening near the click.

    Distance remains important, but textual/paragraph evidence can beat a nearby
    running heading or a deferred-year continuation. X is used as a soft cue so
    the same routine remains safe for multi-column pages.
    """
    same = [(i, ln) for i, ln in enumerate(lines) if int(ln.page) == int(page)]
    if not same:
        return None, None, None
    radius = max(10.0, float(margin_pt))
    candidates = []
    x_aligned_exists = False
    if x is not None:
        for _, ln in same:
            cy = (float(ln.y0) + float(ln.y1)) / 2.0
            if abs(cy - float(y)) > radius * 1.35:
                continue
            lx0, lx1 = float(getattr(ln, "x0", 0.0)), float(getattr(ln, "x1", 0.0))
            if lx0 - 18.0 <= float(x) <= lx1 + 18.0:
                x_aligned_exists = True
                break
    for i, ln in same:
        cy = (float(ln.y0) + float(ln.y1)) / 2.0
        dy = abs(cy - float(y))
        if dy > radius * 1.35:
            continue
        opening = _author_opening_score(lines, i)
        x_bonus = 0.0
        if x is not None:
            lx0, lx1 = float(getattr(ln, "x0", 0.0)), float(getattr(ln, "x1", 0.0))
            aligned = lx0 - 18.0 <= float(x) <= lx1 + 18.0
            if x_aligned_exists and not aligned:
                continue
            x_bonus = 1.0 if aligned else -1.0
        # Plausibility is deliberately stronger than a few points of vertical distance.
        total = opening * 3.0 + x_bonus - (dy / max(radius, 1.0)) * 2.2
        candidates.append((total, opening, -dy, i, ln))
    if candidates:
        candidates.sort(reverse=True, key=lambda t: (t[0], t[1], t[2]))
        best = candidates[0]
        # Only override the geometrically nearest line when there is real start evidence.
        if best[1] >= 1.0:
            return best[3], best[4], abs(((float(best[4].y0)+float(best[4].y1))/2.0)-float(y))
    i, ln = min(same, key=lambda pair: abs(((float(pair[1].y0)+float(pair[1].y1))/2.0)-float(y)))
    return i, ln, abs(((float(ln.y0)+float(ln.y1))/2.0)-float(y))


def _trusted_click_prefix(candidate_lines, clicked_page: int, clicked_y: float, first_record: dict, known_records: list[dict[str, Any]] | None = None) -> tuple[str, list] | tuple[None, list]:
    """Recover the reference beginning at an explicit human click.

    The click is stronger evidence than the automatic segmenter.  Besides the
    segmenter's first returned record, use the already-visible reference list as
    a set of *known later boundaries*.  This matters for the first reference of
    RIEA: even if a local/manual segmentation starts at the deferred-year line
    or jumps to Altieri, the earliest known boundary after the click is Altieri,
    so every physical line from Alonso up to (but not including) Altieri is
    recovered as one candidate.
    """
    if not candidate_lines or not first_record:
        return None, []
    same_page = [(i, ln) for i, ln in enumerate(candidate_lines) if int(ln.page) == int(clicked_page)]
    if not same_page:
        return None, []
    click_i, _ = min(same_page, key=lambda pair: abs(((float(pair[1].y0)+float(pair[1].y1))/2.0)-float(clicked_y)))

    known_texts: list[str] = []
    for record in known_records or []:
        text = _norm(record.get("reference_text") or record.get("text") or "")
        if text and text not in known_texts:
            known_texts.append(text)

    # Prefer a boundary already present in the thesis list. A local segmenter can
    # mistakenly treat the deferred-year continuation (e.g. "G. (2002)") as
    # the first record; using that as the stop would return only the author line.
    boundary_i = None
    if known_texts:
        for i in range(click_i, len(candidate_lines)):
            if any(_line_opens_record(candidate_lines[i].text, text) for text in known_texts):
                boundary_i = i
                break
    if boundary_i is None:
        first_text = _norm(first_record.get("reference_text") or first_record.get("text") or "")
        if first_text:
            for i in range(click_i, len(candidate_lines)):
                if _line_opens_record(candidate_lines[i].text, first_text):
                    boundary_i = i
                    break
    if boundary_i is None or boundary_i <= click_i:
        return None, []
    prefix_lines = candidate_lines[click_i:boundary_i]
    text = _norm(" ".join(ln.text for ln in prefix_lines))
    return (text or None), prefix_lines


@app.post("/api/documents/{doc_id}/references/candidate-from-point")
def reference_candidate_from_point(doc_id: str, payload: ReferenceCandidatePayload):
    """Use a user click as a trusted start hint and let the real segmenter recover one reference."""
    doc = STORE.get(doc_id)
    page = int(payload.page)
    if not 1 <= page <= int(doc.get("page_count") or 0):
        raise HTTPException(400, "Página fuera del PDF")
    try:
        lines = extract_layout_lines_range(doc["path"], page, page)
        if not lines:
            raise ValueError("La página no contiene texto extraíble")
        margin_pt = max(0.0, min(float(payload.margin_pt or PDF_CLICK_MARGIN_PT), 72.0))
        start_idx, nearest, snap_distance = _select_click_start_line(lines, page, float(payload.y), payload.x, margin_pt)
        if nearest is None:
            raise ValueError("No se pudo ubicar una línea cercana al clic")
        # Keep physical context above the snapped opening. The selected line itself
        # remains the trusted human start used below; the extra margin is only for
        # the local segmenter, never a request to include preceding text.
        start_y = max(0.0, float(nearest.y0) - margin_pt)
        end_page = int(payload.end_page or page)
        if end_page < page:
            end_page = page
        spec = ManualZoneSpec(
            start_page=page,
            start_y=start_y,
            end_page=min(end_page, int(doc["page_count"])),
            end_y=payload.end_y if end_page == int(payload.end_page or end_page) else None,
            heading_text=payload.section_label or "Referencias",
            note="candidate_from_pdf_click",
        )
        candidate_lines = extract_layout_lines_range(doc["path"], page, spec.end_page)
        processed = process_manual_zone(doc["path"], spec, lines=candidate_lines)
    except Exception as exc:
        raise HTTPException(400, str(exc))
    records = list(processed.get("records") or [])
    if not records:
        raise HTTPException(404, "No se pudo formar una referencia desde ese punto")
    first = records[0]
    text = _norm(first.get("reference_text") or first.get("text") or "")
    if not text:
        raise HTTPException(404, "El tramo seleccionado no produjo texto bibliográfico")

    visible = [r for r in doc.get("references", []) if not r.get("deleted")]
    # A human click is stronger evidence than the automatic first-boundary rule.
    # Known references already present in the thesis also act as reliable later
    # boundaries, so a missing first entry is recovered even when the local
    # segmenter starts too late.
    trusted_text, trusted_lines = _trusted_click_prefix(candidate_lines, page, float(nearest.y0), first, visible)
    trusted_start = bool(trusted_text)
    if trusted_text:
        text = trusted_text

    duplicate = None
    best_dup = 0.0
    for r in visible:
        score = _reference_match_score(text, r.get("text") or "")
        if score > best_dup:
            best_dup, duplicate = score, r
    if best_dup < 0.88:
        duplicate = None

    before_ref = None
    following_match = None
    following_records = records[:8] if trusted_start else records[1:8]
    for later in following_records:
        later_text = _norm(later.get("reference_text") or later.get("text") or "")
        if not later_text:
            continue
        best = (0.0, None)
        for r in visible:
            score = _reference_match_score(later_text, r.get("text") or "")
            if score > best[0]:
                best = (score, r)
        if best[0] >= 0.82 and best[1] is not None:
            before_ref = best[1]
            following_match = {"reference_id": best[1]["id"], "score": round(best[0], 3), "text": best[1].get("text", "")}
            break

    if trusted_start and trusted_lines:
        out_start_page = int(trusted_lines[0].page)
        out_end_page = int(trusted_lines[-1].page)
    else:
        out_start_page = int(first.get("reference_start_page") or page)
        out_end_page = int(first.get("reference_end_page") or first.get("reference_start_page") or page)
    return {
        "text": text,
        "start_page": out_start_page,
        "end_page": out_end_page,
        "start_y": round(float(nearest.y0), 2),
        "clicked_line": nearest.text,
        "raw_click_y": round(float(payload.y), 2),
        "raw_click_x": (round(float(payload.x), 2) if payload.x is not None else None),
        "snap_distance_pt": round(float(snap_distance or 0.0), 2),
        "selection_margin_pt": round(margin_pt, 1),
        "trusted_click_override": trusted_start,
        "before_ref_id": before_ref.get("id") if before_ref else None,
        "before_ref_order": before_ref.get("order") if before_ref else None,
        "following_match": following_match,
        "already_present_ref_id": duplicate.get("id") if duplicate else None,
        "already_present_score": round(best_dup, 3) if duplicate else None,
        "source": "human_pdf_click",
    }


def _get_ref(doc: dict[str, Any], ref_id: str) -> dict[str, Any]:
    for r in doc.get("references", []):
        if r.get("id") == ref_id and not r.get("deleted"):
            return r
    raise HTTPException(404, "Referencia no encontrada")


@app.get("/api/documents/{doc_id}/references/{ref_id}/origin")
def reference_origin(doc_id: str, ref_id: str):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    start = int(ref.get("start_page") or 1)
    end = int(ref.get("end_page") or start)
    candidate_pages = list(range(max(1, start), min(doc["page_count"], end) + 1))[:4]
    full_text = _norm(ref.get("original_text") or ref.get("text"))
    needle = full_text[:120].lower()
    first_tokens = " ".join(needle.split()[:7])
    candidates = []
    # Lines in reading order across the reference's own pages, so the span can be walked forward.
    ordered: list[tuple[int, Any, float, float]] = []
    for page in candidate_pages:
        lines = extract_layout_lines_range(doc["path"], page, page)
        width, height = page_dimensions(doc["path"], page)
        scored = []
        for line in lines:
            ordered.append((page, line, width, height))
            text = _norm(line.text).lower()
            score = 0
            if first_tokens and first_tokens in text:
                score += 10
            toks = set(first_tokens.split())
            if toks:
                score += sum(t in text for t in toks) / len(toks)
            if score:
                scored.append((score, line))
        for score, line in sorted(scored, key=lambda x: x[0], reverse=True)[:3]:
            candidates.append({
                "page": page, "score": round(float(score), 3),
                "x0": line.x0, "y0": line.y0, "x1": line.x1, "y1": line.y1,
                "page_width": width, "page_height": height, "text": line.text,
            })
    span = _reference_span(ordered, first_tokens, len(full_text))
    return {"reference_id": ref_id, "start_page": start, "end_page": end,
            "candidates": candidates[:5], "span": span}


_SPANS_CACHE: dict[tuple[Any, ...], dict[str, Any]] = {}


@app.get("/api/documents/{doc_id}/reference-spans")
def reference_spans(doc_id: str, start_page: int = 1, end_page: int = 0):
    """Every reference's span in one pass, so the viewer never asks page by page."""
    doc = STORE.get(doc_id)
    page_count = int(doc["page_count"])
    first = max(1, int(start_page))
    last = min(page_count, int(end_page) or page_count)
    if last < first:
        raise HTTPException(422, "El rango de páginas no es válido.")
    last = min(last, first + 59)
    # A long section costs ~2 s to parse. Any edit that moves a reference changes this fingerprint.
    fingerprint = hash(tuple(
        (r.get("id"), r.get("start_page"), r.get("end_page"), len(r.get("original_text") or r.get("text") or ""))
        for r in doc.get("references", []) if not r.get("deleted")
    ))
    cache_key = (doc_id, doc.get("path"), first, last, fingerprint)
    if cache_key in _SPANS_CACHE:
        return _SPANS_CACHE[cache_key]
    result = _compute_reference_spans(doc_id, doc, first, last, page_count)
    if len(_SPANS_CACHE) >= 32:
        _SPANS_CACHE.pop(next(iter(_SPANS_CACHE)))
    _SPANS_CACHE[cache_key] = result
    return result


def _compute_reference_spans(doc_id: str, doc: dict[str, Any], first: int, last: int, page_count: int) -> dict[str, Any]:
    # One parse per page feeds every reference on it; per-reference lookups would reparse it each time.
    lines_by_page: dict[int, tuple[list[Any], float, float]] = {}
    for page in range(first, last + 1):
        width, height = page_dimensions(doc["path"], page)
        lines_by_page[page] = (extract_layout_lines_range(doc["path"], page, page), width, height)
    spans: dict[str, list[dict[str, Any]]] = {}
    # References are located in reading order: each one starts after the previous one, so
    # «Ward (2008a)», «(2008b)», «(2008c)» with identical authors do not all land on the first line.
    cursor = None
    refs = sorted((r for r in doc.get("references", []) if not r.get("deleted")),
                  key=lambda r: (int(r.get("start_page") or 0), int(r.get("order") or 0)))
    for ref in refs:
        start = int(ref.get("start_page") or 1)
        if start < first or start > last:
            continue
        span = _span_for_reference(ref, lines_by_page, page_count, after=cursor)
        if span:
            spans[ref["id"]] = span
            cursor = (span[0]["page"], span[0]["y0"])
    return {"document_id": doc_id, "start_page": first, "end_page": last, "spans": spans}


def _span_for_reference(ref: dict[str, Any], lines_by_page, page_count: int, after=None) -> list[dict[str, Any]]:
    start = int(ref.get("start_page") or 1)
    end = int(ref.get("end_page") or start)
    ordered: list[tuple[int, Any, float, float]] = []
    for page in list(range(max(1, start), min(page_count, end) + 1))[:4]:
        entry = lines_by_page.get(page)
        if not entry:
            continue
        lines, width, height = entry
        ordered.extend((page, line, width, height) for line in lines)
    full_text = _norm(ref.get("original_text") or ref.get("text"))
    first_tokens = " ".join(full_text[:120].lower().split()[:7])
    return _reference_span(ordered, first_tokens, len(full_text), full_text=full_text, after=after)


_DOCUMENT_MAP_CACHE: dict[tuple[str, Any], dict[str, Any]] = {}
TOC_BIBLIOGRAPHY_RX = re.compile(
    r"^\s*((?:bibliograf[ií]a|referencias(?:\s+bibliogr[aá]ficas)?|obras\s+citadas|fuentes)[^\n]*?)[\s.…·_\-]*(\d{1,4})\s*$",
    re.I,
)
MAP_YEAR_RX = re.compile(r"\b(?:1[5-9]\d{2}|20\d{2})[a-z]?\b")
MAP_SOURCE_RX = re.compile(
    r"\b(?:editorial|ediciones|eds?\.|pp?\.|p[aá]gs?\.|vol\.|revista|journal|press|university|universidad|"
    r"doi|isbn|madrid|barcelona|buenos\s+aires|m[eé]xico|paid[oó]s|alianza|siglo\s+xxi|amorrortu|"
    r"fondo\s+de\s+cultura)\b|https?://",
    re.I,
)
MAP_AUTHOR_RX = re.compile(r"^[\W\d]*[A-ZÁÉÍÓÚÑ][\w'’\-]+(?:\s[A-ZÁÉÍÓÚÑ][\w'’\-]+)?,\s")
MAP_DEFINITION_RX = re.compile(r"^\s*\d+\.\s+[a-z]{1,4}\.\s")


def _bibliographic_evidence(lines) -> dict[str, Any]:
    """Structure alone (bullets, numbers) also matches questions, definitions and conclusions."""
    texts = [_norm(line.text) for line in lines if _norm(line.text)]
    dated = sum(bool(MAP_YEAR_RX.search(t)) for t in texts)
    sourced = sum(bool(MAP_SOURCE_RX.search(t)) for t in texts)
    authored = sum(bool(MAP_AUTHOR_RX.match(t)) for t in texts)
    questions = sum("¿" in t or t.endswith("?") for t in texts)
    definitions = sum(bool(MAP_DEFINITION_RX.match(t)) for t in texts)
    accepted = (
        dated >= 3
        and sourced + authored >= 3
        and questions * 3 < max(1, len(texts))
        and definitions * 4 < max(1, len(texts))
    )
    return {"accepted": accepted, "lines": len(texts), "with_year": dated,
            "with_source": sourced, "with_author": authored,
            "questions": questions, "definitions": definitions}


@app.get("/api/documents/{doc_id}/document-map")
def document_map(doc_id: str):
    """Where bibliography could be in the whole document, beyond the sections already detected."""
    from biblioref.audit.detection import detect_reference_like_regions

    doc = STORE.get(doc_id)
    key = (doc_id, doc.get("updated_at"))
    if key in _DOCUMENT_MAP_CACHE:
        return _DOCUMENT_MAP_CACHE[key]
    page_count = int(doc["page_count"])
    lines = extract_layout_lines_range(doc["path"], 1, page_count)
    sections = []
    covered: set[int] = set()
    for sec in doc.get("sections") or []:
        start = int(sec.get("start_page") or 1)
        end = int(sec.get("end_page") or start)
        sections.append({"label": sec.get("label") or "Referencias", "start_page": start, "end_page": end})
        covered.update(range(start, end + 1))
    toc = []
    for line in lines:
        if line.page > min(20, page_count):
            continue
        match = TOC_BIBLIOGRAPHY_RX.match(_norm(line.text))
        if not match:
            continue
        target = int(match.group(2))
        if 1 <= target <= page_count and target != line.page:
            toc.append({"page": int(line.page), "label": match.group(1).strip(" .…·_-"), "target_page": target})
    candidates, discarded = [], []
    for region in detect_reference_like_regions(lines, page_count):
        start, end = int(region["start_page"]), int(region["end_page"])
        if all(page in covered for page in range(start, end + 1)):
            continue
        evidence = _bibliographic_evidence([line for line in lines if start <= line.page <= end])
        entry = {**region, "bibliographic_evidence": evidence}
        (candidates if evidence["accepted"] else discarded).append(entry)
    result = {"document_id": doc_id, "page_count": page_count, "sections": sections,
              "toc": toc, "candidates": candidates, "discarded": discarded}
    _DOCUMENT_MAP_CACHE.clear()
    _DOCUMENT_MAP_CACHE[key] = result
    return result


def _reference_span(ordered, first_tokens: str, target_len: int, full_text: str = "", after=None) -> list[dict[str, Any]]:
    """Walk forward from the line that starts the reference until its text is consumed.

    Scoring alone marks neighbouring references too, so the box must follow the
    reference's own length instead of every line that merely looks similar.
    ``after`` (page, y) skips lines before the previous reference, and the whole
    first line is compared with the reference's opening so near-identical
    neighbours (same authors, 2008a/2008b) are told apart.
    """
    if not ordered or not first_tokens:
        return []
    tokens = set(first_tokens.split())
    opening = re.sub(r"\s+", "", (full_text or "").lower())

    def pick(skip_before):
        best, best_score = None, 0.0
        for index, (page, line, _w, _h) in enumerate(ordered):
            if skip_before and (page < skip_before[0] or (page == skip_before[0] and line.y0 <= skip_before[1] + 0.5)):
                continue
            text = _norm(line.text).lower()
            score = 10.0 if first_tokens in text else 0.0
            if tokens:
                score += sum(t in text for t in tokens) / len(tokens)
            if opening and score:
                compact = re.sub(r"\s+", "", text)
                shared = len(os.path.commonprefix([compact, opening]))
                if shared >= 15:
                    score += 20.0 * shared / max(len(compact), 1)
            # Earlier lines win ties: a reference starts at its first matching line.
            if score > best_score:
                best, best_score = index, score
        return best

    anchor = pick(after) if after else None
    if anchor is None:
        anchor = pick(None)
    if anchor is None:
        return []
    boxes: dict[int, dict[str, Any]] = {}
    consumed = 0
    for page, line, width, height in ordered[anchor:anchor + 14]:
        box = boxes.setdefault(page, {
            "page": page, "x0": line.x0, "y0": line.y0, "x1": line.x1, "y1": line.y1,
            "page_width": width, "page_height": height,
        })
        box["x0"] = min(box["x0"], line.x0)
        box["y0"] = min(box["y0"], line.y0)
        box["x1"] = max(box["x1"], line.x1)
        box["y1"] = max(box["y1"], line.y1)
        consumed += len(_norm(line.text)) + 1
        if consumed >= target_len * 0.92:
            break
    return list(boxes.values())


@app.patch("/api/documents/{doc_id}/references/{ref_id}")
def edit_reference(doc_id: str, ref_id: str, payload: RefEditPayload):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    old = ref["text"]
    old_review = deepcopy(ref.get("parsed_review") or {})
    ref["text"] = _norm(payload.text)
    ref.pop("parsed_review", None)
    _reset_reference_resolution(ref, preserve_human=False)
    _refresh_reference_parse(ref)
    ref["edited"] = True
    if payload.status:
        ref["status"] = payload.status
    _log_correction(doc, "edit_reference", {"reference_id": ref_id, "before": old, "after": ref["text"], "parsed_review_reset": bool((old_review.get("overrides") or {}))})
    doc["updated_at"] = _now()
    STORE.save()
    return _reference_with_parse(ref)


@app.patch("/api/documents/{doc_id}/references/{ref_id}/parsed-review")
def edit_reference_parsed_review(doc_id: str, ref_id: str, payload: ParsedReviewPayload):
    """Store field corrections as a sidecar; parser output and raw text stay immutable."""
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    requested = set(payload.overrides) | set(payload.unset_fields)
    unknown = sorted(requested - set(PARSED_REVIEW_FIELDS))
    if unknown:
        raise HTTPException(400, "Campos de revisión no admitidos: " + ", ".join(unknown))

    review = ref.setdefault("parsed_review", {})
    overrides = review.setdefault("overrides", {})
    before = deepcopy(overrides)
    for field in payload.unset_fields:
        overrides.pop(field, None)
    for field, value in payload.overrides.items():
        if isinstance(value, (dict, list, tuple, set)):
            raise HTTPException(400, f"El campo {field} debe ser un valor escalar")
        overrides[field] = _norm(value)

    changed_fields = sorted({key for key in set(before) | set(overrides) if before.get(key) != overrides.get(key)})
    if changed_fields:
        review["status"] = "edited" if overrides else "parser_original"
        review["updated_at"] = _now()
        review["source"] = "human_spreadsheet"
        _reset_reference_resolution(ref, preserve_human=True)
        _log_correction(doc, "edit_parsed_fields", {
            "reference_id": ref_id,
            "fields": changed_fields,
            "before_overrides": before,
            "after_overrides": deepcopy(overrides),
            "raw_reference_unchanged": True,
            "parser_output_unchanged": True,
        })
        doc["updated_at"] = _now()
        STORE.save()
    return _reference_with_parse(ref)


@app.post("/api/documents/{doc_id}/references/{ref_id}/resolve")
def resolve_reference_layers(doc_id: str, ref_id: str, payload: ResolutionRunPayload):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    row = _reference_with_parse(ref)
    runner = InteractiveResolutionRunner()
    try:
        result = runner.run(
            raw_reference=row.get("text") or "",
            parsed=row.get("effective_parsed") or row.get("parsed") or {},
            layers=payload.layers,
        ).as_dict()
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    for key in RESOLUTION_SIDECAR_KEYS:
        ref.pop(key, None)
    ref["resolver_query"] = result["query"]
    ref["resolver_batch"] = result["batch"]
    ref["resolver_candidates"] = result["candidates"]
    if result.get("consensus") is not None:
        ref["consensus"] = result["consensus"]
    if result.get("semantic") is not None:
        ref["semantic"] = result["semantic"]
    if result.get("rag") is not None:
        ref["rag_result"] = result["rag"]
    if result.get("adjudication") is not None:
        adjudication = result["adjudication"]
        if adjudication.get("rag_version") is not None:
            ref["rag_adjudication"] = adjudication
        else:
            ref["adjudication"] = adjudication
    if result.get("decision") is not None:
        ref["resolution_decision"] = result["decision"]
    ref["resolution_run"] = {
        "version": result["version"],
        "requested_layers": result["requested_layers"],
        "executed_layers": result["executed_layers"],
        "layer_errors": result["layer_errors"],
        "timestamp": _now(),
        "query_source": "effective_parsed" if row.get("parsed_review", {}).get("overrides") else "parser",
    }
    doc["updated_at"] = _now()
    STORE.save()
    return _reference_with_parse(ref)


@app.post("/api/documents/{doc_id}/references/{ref_id}/human-resolution")
def confirm_reference_resolution(doc_id: str, ref_id: str, payload: HumanResolutionPayload):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    decision = _norm(payload.decision).casefold()
    if decision not in {"candidate", "none"}:
        raise HTTPException(400, "decision debe ser candidate o none")

    if decision == "none":
        ref["human_resolution"] = {
            "status": "human_confirmed_none",
            "candidate_index": None,
            "note": _norm(payload.note),
            "timestamp": _now(),
        }
        _log_correction(doc, "human_resolution_confirmation", {
            "reference_id": ref_id, "decision": "none", "note": _norm(payload.note),
        })
    else:
        candidates = list(ref.get("resolver_candidates") or [])
        index = payload.candidate_index
        if isinstance(index, bool) or not isinstance(index, int) or index < 0 or index >= len(candidates):
            raise HTTPException(400, "candidate_index fuera de rango")
        data = candidates[index]
        identity_key = _norm(data.get("identity_key")) or None
        ref["human_resolution"] = {
            "status": "human_confirmed",
            "candidate_index": index,
            "identity_key": identity_key,
            "candidate": deepcopy(data),
            "note": _norm(payload.note),
            "timestamp": _now(),
        }
        _log_correction(doc, "human_resolution_confirmation", {
            "reference_id": ref_id,
            "decision": "candidate",
            "candidate_index": index,
            "identity_key": identity_key,
            "note": _norm(payload.note),
        })
    doc["updated_at"] = _now()
    STORE.save()
    return _reference_with_parse(ref)


@app.delete("/api/documents/{doc_id}/references/{ref_id}")
def remove_reference(doc_id: str, ref_id: str):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    ref["deleted"] = True
    _log_correction(doc, "delete_reference", {"reference_id": ref_id, "text": ref["text"]})
    _renumber(doc)
    STORE.save()
    return _public_doc(doc, True)


def _apply_split(doc: dict[str, Any], ref: dict[str, Any], parts: list[str]) -> list[dict[str, Any]]:
    refs = doc["references"]
    idx = refs.index(ref)
    old = deepcopy(ref)
    new_refs = []
    for part in parts:
        nr = deepcopy(old)
        nr["id"] = f"ref_{uuid.uuid4().hex[:12]}"
        nr["text"] = part
        nr.pop("parsed_review", None)
        _reset_reference_resolution(nr, preserve_human=False)
        _refresh_reference_parse(nr)
        # Each part is located on its own lines: keeping the whole merged original
        # made both halves claim the same block (and the second one find none).
        nr["original_text"] = part
        nr["source"] = "human_split"
        nr["edited"] = True
        nr["status"] = "reviewed"
        new_refs.append(nr)
    refs[idx:idx+1] = new_refs
    _log_correction(doc, "split_reference", {"reference_id": ref.get("id"), "before": old["text"], "parts": parts})
    _renumber(doc)
    return new_refs


@app.post("/api/documents/{doc_id}/references/{ref_id}/split")
def split_reference(doc_id: str, ref_id: str, payload: SplitPayload):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    parts = [_norm(p) for p in payload.parts if _norm(p)]
    if len(parts) < 2:
        raise HTTPException(400, "La división necesita al menos dos fragmentos no vacíos")
    _apply_split(doc, ref, parts)
    STORE.save()
    return _public_doc(doc, True)


@app.post("/api/documents/{doc_id}/references/{ref_id}/merge")
def merge_reference(doc_id: str, ref_id: str, payload: MergePayload):
    doc = STORE.get(doc_id)
    a = _get_ref(doc, ref_id)
    b = _get_ref(doc, payload.with_ref_id)
    refs = doc["references"]
    ia, ib = refs.index(a), refs.index(b)
    if abs(ia - ib) != 1:
        raise HTTPException(400, "Sólo se pueden unir referencias contiguas")
    first, second = (a, b) if ia < ib else (b, a)
    merged = deepcopy(first)
    merged["id"] = f"ref_{uuid.uuid4().hex[:12]}"
    merged["text"] = _norm(first["text"] + " " + second["text"])
    merged.pop("parsed_review", None)
    _reset_reference_resolution(merged, preserve_human=False)
    _refresh_reference_parse(merged)
    merged["original_text"] = _norm(first.get("original_text", "") + " " + second.get("original_text", ""))
    merged["start_page"] = first.get("start_page")
    merged["end_page"] = second.get("end_page") or first.get("end_page")
    merged["source"] = "human_merge"
    merged["edited"] = True
    merged["status"] = "reviewed"
    start = min(ia, ib)
    refs[start:start+2] = [merged]
    _log_correction(doc, "merge_references", {"reference_ids": [a["id"], b["id"]], "after": merged["text"]})
    _renumber(doc)
    STORE.save()
    return _public_doc(doc, True)


@app.post("/api/documents/{doc_id}/references")
def insert_reference(doc_id: str, payload: InsertPayload):
    doc = STORE.get(doc_id)
    if not _norm(payload.text):
        raise HTTPException(400, "La referencia no puede quedar vacía")
    source = payload.source if payload.source in {"human_insert", "human_pdf_click", "human_paste"} else "human_insert"
    ref = {
        "id": f"ref_{uuid.uuid4().hex[:12]}", "order": 0,
        "text": _norm(payload.text),
        "original_text": _norm(payload.text) if source == "human_pdf_click" else "",
        "start_page": payload.start_page, "end_page": payload.end_page or payload.start_page,
        "start_y": payload.start_y, "end_y": payload.end_y,
        "start_x": None, "zone_id": None, "style": None, "boundary_evidence": "",
        "audit_flags": "", "status": "reviewed", "source": source, "edited": True, "deleted": False,
    }
    _refresh_reference_parse(ref)
    refs = doc.setdefault("references", [])
    if payload.before_ref_id:
        target = _get_ref(doc, payload.before_ref_id)
        refs.insert(refs.index(target), ref)
    elif payload.after_ref_id:
        target = _get_ref(doc, payload.after_ref_id)
        refs.insert(refs.index(target) + 1, ref)
    else:
        refs.append(ref)
    _log_correction(doc, "insert_reference", {
        "reference_id": ref["id"], "text": ref["text"], "page": ref["start_page"],
        "start_y": ref.get("start_y"), "source": source, "before_ref_id": payload.before_ref_id,
        "after_ref_id": payload.after_ref_id,
    })
    _renumber(doc)
    STORE.save()
    return _public_doc(doc, True)


@app.post("/api/documents/{doc_id}/references/{ref_id}/mark")
def mark_reference(doc_id: str, ref_id: str, status: str = "reviewed"):
    doc = STORE.get(doc_id)
    ref = _get_ref(doc, ref_id)
    ref["status"] = status
    _log_correction(doc, "mark_reference", {"reference_id": ref_id, "status": status})
    STORE.save()
    return ref


def _visible_refs(doc: dict[str, Any]):
    # «------------- (2013)» takes the author of the explicit entry above it.
    return inherit_repeated_authors(
        [_reference_with_parse(r) for r in doc.get("references", []) if not r.get("deleted")]
    )


@app.get("/api/documents/{doc_id}/audit-checks")
def audit_checks(doc_id: str):
    doc = STORE.get(doc_id)
    coverage = deepcopy((doc.get("analysis") or {}).get("coverage") or {})
    decisions = deepcopy(doc.get("audit_decisions") or {})
    for item in coverage.get("possible_missing", []) or []:
        if item.get("id") in decisions:
            item["human_decision"] = decisions[item["id"]]
    for item in coverage.get("page_boundaries", []) or []:
        if item.get("id") in decisions:
            item["human_decision"] = decisions[item["id"]]
    return {"coverage": coverage, "decisions": decisions}


@app.post("/api/documents/{doc_id}/audit-checks/{check_id}/decision")
def audit_check_decision(doc_id: str, check_id: str, payload: AuditDecisionPayload):
    doc = STORE.get(doc_id)
    allowed = {"split", "continue", "ignore", "pending", "accepted_missing", "resolved"}
    if payload.decision not in allowed:
        raise HTTPException(400, "Decisión de control no válida")
    entry = {"decision": payload.decision, "note": payload.note, "timestamp": _now()}
    doc.setdefault("audit_decisions", {})[check_id] = entry
    _log_correction(doc, "audit_check_decision", {"check_id": check_id, **entry})
    doc["updated_at"] = _now()
    STORE.save()
    return {"ok": True, "check_id": check_id, "human_decision": entry}


MISSING_CLOSED = {"ignore", "resolved", "accepted_missing"}
_BAND_ABOVE = 34.0   # contexto arriba del párrafo, en puntos
_BAND_BELOW = 26.0   # y abajo: sin vecinos a la vista no se sabe dónde cae el recorte
_BAND_LINE = 15.5


def _doc_title(doc: dict[str, Any]) -> str:
    return _norm(doc.get("title") or (doc.get("metadata") or {}).get("title") or doc.get("filename") or "")


def _open_missing(doc: dict[str, Any]) -> list[dict[str, Any]]:
    """Possible missing references of a thesis that nobody has decided on yet."""
    decisions = doc.get("audit_decisions") or {}
    coverage = (doc.get("analysis") or {}).get("coverage") or {}
    return [x for x in coverage.get("possible_missing") or []
            if (decisions.get(x.get("id")) or {}).get("decision") not in MISSING_CLOSED]


def _missing_paragraph_rect(doc: dict[str, Any], item: dict[str, Any]) -> tuple[float, float]:
    """Vertical extent of the paragraph, read off the page when the lines can be parsed."""
    y = float(item.get("y") or 0)
    text = _norm(item.get("text"))
    first = _norm(item.get("first_line"))
    page = int(item.get("page") or 1)
    try:
        lines = [l for l in extract_layout_lines_range(doc["path"], page, page) if float(l.y1) > y - 2]
        lines.sort(key=lambda l: float(l.y0))
        consumed, bottom = 0, None
        for line in lines[:8]:
            bottom = float(line.y1)
            consumed += len(_norm(line.text)) + 1
            if consumed >= max(1, len(text)) * 0.92:
                break
        if bottom and bottom > y:
            return y, bottom
    except Exception:
        pass
    per_line = max(30, len(first) or 90)
    return y, y + _BAND_LINE * max(1, min(6, round(len(text) / per_line)))


_REFERENCE_FLAT_CACHE: dict[str, tuple[Any, list[tuple[dict[str, Any], str, list[int]]]]] = {}
MISSING_NEEDLE = 70


def _flatten_with_map(text: str) -> tuple[str, list[int]]:
    """Letters and digits only, plus where each one started in the original text."""
    chars: list[str] = []
    positions: list[int] = []
    for index, char in enumerate(text or ""):
        for piece in unicodedata.normalize("NFKD", char):
            if unicodedata.combining(piece) or not piece.isalnum():
                continue
            chars.append(piece.lower())
            positions.append(index)
    return "".join(chars), positions


def _reference_flat_index(doc: dict[str, Any]):
    key = doc.get("id")
    stamp = (doc.get("updated_at"), len(doc.get("references") or []))
    cached = _REFERENCE_FLAT_CACHE.get(key)
    if cached and cached[0] == stamp:
        return cached[1]
    index = []
    for ref in _visible_refs(doc):
        flat, positions = _flatten_with_map(_norm(ref.get("text")))
        index.append((ref, flat, positions))
    if len(_REFERENCE_FLAT_CACHE) >= 24:
        _REFERENCE_FLAT_CACHE.pop(next(iter(_REFERENCE_FLAT_CACHE)))
    _REFERENCE_FLAT_CACHE[key] = (stamp, index)
    return index


def _missing_inside_reference(doc: dict[str, Any], item: dict[str, Any]) -> dict[str, Any] | None:
    """¿Este «faltante» ya está adentro de una referencia extraída?

    Pasa cuando dos referencias quedaron pegadas: el párrafo está en el PDF y también
    dentro del texto de la referencia anterior, así que extraerlo lo duplicaría.

    También pasa, y es el caso más común, cuando el párrafo era la segunda línea de
    esa misma referencia. Ahí no hay nada que dividir: ``cut_at`` queda en ``None``
    y el aviso deja de pedir una decisión que partiría la referencia en dos.
    """
    needle = _flatten_with_map(_norm(item.get("first_line")))[0][:MISSING_NEEDLE]
    if len(needle) < 25:
        return None
    for ref, flat, positions in _reference_flat_index(doc):
        at = flat.find(needle)
        if at < 0:
            continue
        text = _norm(ref.get("text"))
        cut_at = positions[at] if at > 0 else None
        continuation = cut_at is not None and not cut_opens_a_record(text[cut_at:])
        return {
            "reference_id": ref.get("id"),
            "order": ref.get("order"),
            "length": len(text),
            "cut_at": None if continuation else cut_at,
            "already_extracted": at == 0,
            "continuation": continuation,
        }
    return None


_PREVIEW_MIN_RUN = 10  # letras seguidas para dar por encontrado un tramo del recuadro


def _missing_preview(doc: dict[str, Any], item: dict[str, Any], inside: dict[str, Any] | None) -> dict[str, Any] | None:
    """Cómo quedan las referencias que tocan el recuadro, y qué parte no quedó en ninguna.

    El recuadro suele abarcar el final de una referencia y el comienzo de la siguiente,
    así que se muestran las dos: en cada una, los tramos que vienen del recuadro; aparte,
    lo que está en el recuadro pero en ninguna referencia (un renglón perdido).
    """
    if not inside:
        return None
    index = _reference_flat_index(doc)
    at = next((i for i, row in enumerate(index) if row[0].get("id") == inside.get("reference_id")), None)
    if at is None:
        return None
    candidates = index[at:at + 2]
    item_text = _norm(item.get("text"))
    flat, positions = _flatten_with_map(item_text)
    spans: dict[str, list[list[int]]] = {row[0].get("id"): [] for row in candidates}
    lost: list[list[int]] = []
    i = 0
    while i < len(flat):
        best_len, best_row, best_at = 0, None, -1
        for row in candidates:
            ref_flat = row[1]
            lo, hi = 0, len(flat) - i
            while lo < hi:  # el tramo más largo desde i que aparece en esta referencia
                mid = (lo + hi + 1) // 2
                if flat[i:i + mid] in ref_flat:
                    lo = mid
                else:
                    hi = mid - 1
            if lo > best_len:
                best_len, best_row, best_at = lo, row, ref_flat.find(flat[i:i + lo])
        if best_len >= _PREVIEW_MIN_RUN:
            ref, _, ref_positions = best_row
            spans[ref.get("id")].append([ref_positions[best_at], ref_positions[best_at + best_len - 1] + 1])
            i += best_len
            continue
        if lost and lost[-1][1] == i:
            lost[-1][1] = i + 1
        else:
            lost.append([i, i + 1])
        i += 1
    leftover = []
    for a, b in lost:
        if b - a < 4:  # un número de página o una sigla suelta no es un renglón perdido
            continue
        end = positions[b - 1] + 1
        while end < len(item_text) and item_text[end] in ".)]»”\"":
            end += 1
        piece = item_text[positions[a]:end].strip(" ,;")
        if piece:
            leftover.append(piece)
    return {
        "refs": [{"id": ref.get("id"), "order": ref.get("order"), "text": _norm(ref.get("text")),
                  "spans": spans.get(ref.get("id")) or []} for ref, _, _ in candidates],
        "leftover": leftover,
    }


def _split_parts_for_missing(doc: dict[str, Any], item: dict[str, Any], force: bool = False) -> tuple[dict[str, Any], list[str]]:
    inside = _missing_inside_reference(doc, item)
    if inside and force and inside.get("cut_at") is None and inside.get("continuation"):
        # La persona dice que el párrafo es una referencia aparte: se corta donde empieza.
        text = _norm(_get_ref(doc, inside["reference_id"]).get("text"))
        flat, positions = _flatten_with_map(text)
        needle = _flatten_with_map(_norm(item.get("first_line")))[0][:MISSING_NEEDLE]
        found = flat.find(needle)
        if found > 0:
            inside = {**inside, "cut_at": positions[found]}
    if not inside or inside.get("cut_at") is None:
        raise HTTPException(400, "Ese párrafo no está adentro de una referencia extraída")
    ref = _get_ref(doc, inside["reference_id"])
    text = _norm(ref.get("text"))
    head, tail = _norm(text[: inside["cut_at"]]), _norm(text[inside["cut_at"]:])
    if not head or not tail:
        raise HTTPException(400, "El corte dejaría una de las dos partes vacía")
    return ref, [head, tail]


@app.post("/api/documents/{doc_id}/missing/{check_id}/split-host")
def split_host_reference(doc_id: str, check_id: str, force: bool = False):
    """Divide la referencia que se tragó este párrafo, justo donde el párrafo empieza.

    ``force`` corta también cuando el párrafo parecía la continuación de esa referencia:
    lo pide quien revisa y dice que es una referencia aparte.
    """
    doc = STORE.get(doc_id)
    coverage = (doc.get("analysis") or {}).get("coverage") or {}
    item = next((x for x in coverage.get("possible_missing") or [] if x.get("id") == check_id), None)
    if not item:
        raise HTTPException(404, "Ese faltante no existe en esta tesis")
    ref, parts = _split_parts_for_missing(doc, item, force=force)
    new_refs = _apply_split(doc, ref, parts)
    doc.setdefault("audit_decisions", {})[check_id] = {
        "decision": "resolved", "note": "split_host_from_missing", "timestamp": _now(),
    }
    _log_correction(doc, "audit_check_decision", {"check_id": check_id, "decision": "resolved", "note": "split_host_from_missing"})
    STORE.save()
    return {"ok": True, "reference_id": ref.get("id"),
            "parts": [{"id": r.get("id"), "order": r.get("order"), "text": r.get("text")} for r in new_refs]}


class MissingAppendPayload(BaseModel):
    reference_id: str
    text: str = Field(min_length=1, max_length=2000)


@app.post("/api/documents/{doc_id}/missing/{check_id}/append")
def append_to_reference_from_missing(doc_id: str, check_id: str, payload: MissingAppendPayload):
    """Agrega al final de una referencia el tramo del recuadro que no quedó en ninguna."""
    doc = STORE.get(doc_id)
    coverage = (doc.get("analysis") or {}).get("coverage") or {}
    if not any(x.get("id") == check_id for x in coverage.get("possible_missing") or []):
        raise HTTPException(404, "Ese faltante no existe en esta tesis")
    ref = _get_ref(doc, payload.reference_id)
    before = ref["text"]
    ref["text"] = _norm(f"{before} {payload.text}")
    ref.pop("parsed_review", None)
    _reset_reference_resolution(ref, preserve_human=False)
    _refresh_reference_parse(ref)
    ref["edited"] = True
    _log_correction(doc, "edit_reference", {"reference_id": ref["id"], "before": before, "after": ref["text"],
                                            "source": "append_from_missing", "check_id": check_id})
    doc.setdefault("audit_decisions", {})[check_id] = {
        "decision": "resolved", "note": "appended_to_reference_from_missing", "timestamp": _now(),
    }
    doc["updated_at"] = _now()
    STORE.save()
    return _reference_with_parse(ref)


@app.get("/api/documents/{doc_id}/missing/{check_id}/band.png")
def missing_band_png(doc_id: str, check_id: str, zoom: float = 1.6):
    """The strip of page a reviewer needs to judge one possible missing reference."""
    doc = STORE.get(doc_id)
    coverage = (doc.get("analysis") or {}).get("coverage") or {}
    item = next((x for x in coverage.get("possible_missing") or [] if x.get("id") == check_id), None)
    if not item:
        raise HTTPException(404, "Ese faltante no existe en esta tesis")
    top, bottom = _missing_paragraph_rect(doc, item)
    try:
        png = render_band_png(
            doc["path"], int(item.get("page") or 1),
            top - _BAND_ABOVE, bottom + _BAND_BELOW,
            zoom=min(max(zoom, 0.6), 2.5), mark=(top - 3.0, bottom + 2.0),
        )
    except Exception as exc:
        raise HTTPException(400, str(exc))
    return Response(content=png, media_type="image/png", headers={"Cache-Control": "no-store"})


def _audit_parsed(ref: dict[str, Any]) -> dict[str, Any]:
    out = dict(ref.get("parsed") or {})
    out.update((ref.get("parsed_review") or {}).get("overrides") or {})
    return out


def _audit_scalar(value: Any) -> str:
    if isinstance(value, list):
        value = "; ".join(str(v) for v in value if v)
    return _norm(value or "") or "—"


def _crop_name(item: dict[str, Any]) -> str:
    return f"recortes/p{int(item.get('page') or 0)}_{_safe_name(str(item.get('id') or 'falta'))}.png"


def _audit_report(doc: dict[str, Any], with_crops: bool) -> str:
    """Lo extraído y lo que falta, intercalado en el orden del PDF, para revisar la extracción."""
    refs = _visible_refs(doc)
    missing = _open_missing(doc)
    coverage = (doc.get("analysis") or {}).get("coverage") or {}
    merged = coverage.get("merged_suspects") or []
    sections = doc.get("sections") or []
    analysis = doc.get("analysis") or {}
    metadata = doc.get("metadata") or {}

    # La posición de cada referencia en la página: sin eso no se puede intercalar con los faltantes.
    spans: dict[str, Any] = {}
    for section in sections:
        first = max(1, int(section.get("start_page") or 1))
        last = min(int(doc["page_count"]), int(section.get("end_page") or first))
        try:
            spans.update(_compute_reference_spans(doc["id"], doc, first, min(last, first + 59), int(doc["page_count"]))["spans"])
        except Exception:
            continue

    rows: list[tuple[int, float, str, dict[str, Any]]] = []
    for ref in refs:
        box = (spans.get(ref["id"]) or [{}])[0]
        rows.append((int(box.get("page") or ref.get("start_page") or 0), float(box.get("y0") or 0), "ref", ref))
    for item in missing:
        rows.append((int(item.get("page") or 0), float(item.get("y") or 0), "missing", item))
    rows.sort(key=lambda row: (row[0], row[1]))

    by_kind: dict[str, Counter] = {"ref": Counter(), "missing": Counter()}
    for page, _y, kind, _payload in rows:
        by_kind[kind][page] += 1

    out = [
        f"# Auditoría de extracción · {doc.get('thesis_id') or doc.get('id')}",
        "",
        f"- **Tesis**: {_audit_scalar(metadata.get('author'))} ({_audit_scalar(metadata.get('year'))}) · {_audit_scalar(_doc_title(doc))}",
        f"- **PDF**: {doc.get('filename')} · {doc.get('page_count')} páginas",
        "- **Secciones de bibliografía**: " + (", ".join(
            f"{section.get('label') or 'sección'} {section.get('start_page')}–{section.get('end_page')}" for section in sections) or "ninguna"),
        f"- **Extracción**: software {analysis.get('software_version')} · algoritmo {analysis.get('algorithm_version')} · método {analysis.get('detection_method')}",
        f"- **Totales**: {len(refs)} referencias extraídas · {len(missing)} posibles faltantes · {len(merged)} posibles fusionadas",
        "",
        "## Referencias por página",
        "",
        "| página | extraídas | posibles faltantes |",
        "| --- | --- | --- |",
    ]
    for page in sorted(set(by_kind["ref"]) | set(by_kind["missing"])):
        out.append(f"| {page} | {by_kind['ref'].get(page, 0)} | {by_kind['missing'].get(page, 0)} |")

    out += ["", "## Bibliografía en el orden del PDF", "",
            "`EXTRAÍDA` = quedó como referencia · `FALTA` = el párrafo está en el PDF y no quedó como referencia.", ""]
    current_page = None
    for page, _y, kind, payload in rows:
        if page != current_page:
            current_page = page
            out += ["", f"### Página {page}", ""]
        if kind == "ref":
            parsed = _audit_parsed(payload)
            out.append(
                f"- **EXTRAÍDA · REF {int(payload.get('order') or 0):03d}** · autor: {_audit_scalar(parsed.get('author'))}"
                f" · año: {_audit_scalar(parsed.get('year'))} · tipo: {_audit_scalar(parsed.get('publication_type'))}"
                f" · origen: {payload.get('source') or 'automatic'}"
            )
            out.append(f"  > {_norm(payload.get('text'))}")
        else:
            crop = f" · recorte: `{_crop_name(payload)}`" if with_crops else ""
            out.append(f"- **FALTA** · parecido con la referencia más cercana: {payload.get('best_existing_match')}{crop}")
            out.append(f"  > {_norm(payload.get('text'))}")

    if merged:
        out += ["", "## Posibles fusionadas", ""]
        for item in merged:
            out.append(f"- pág. {item.get('page')}")
            out.append(f"  > {_norm(item.get('text'))[:400]}")

    return "\n".join(out) + "\n"


@app.get("/api/documents/{doc_id}/audit-export.zip")
def audit_export_zip(doc_id: str, crops: bool = True, max_crops: int = 300):
    """Un .zip con el informe de auditoría de la tesis y el recorte de cada faltante."""
    doc = STORE.get(doc_id)
    thesis = _safe_name(str(doc.get("thesis_id") or doc.get("id")))
    report = _audit_report(doc, crops)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr(f"auditoria_{thesis}.md", report)
        if crops:
            for item in _open_missing(doc)[: max(0, int(max_crops))]:
                top, bottom = _missing_paragraph_rect(doc, item)
                try:
                    png = render_band_png(
                        doc["path"], int(item.get("page") or 1),
                        top - _BAND_ABOVE, bottom + _BAND_BELOW,
                        zoom=1.6, mark=(top - 3.0, bottom + 2.0),
                    )
                except Exception:
                    continue
                bundle.writestr(_crop_name(item), png)
    return Response(
        content=buffer.getvalue(), media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="auditoria_{thesis}.zip"',
                 "Cache-Control": "no-store"},
    )


def _split_all_in_document(doc: dict[str, Any]) -> tuple[int, int]:
    """Separa, de una, todas las pegadas de una tesis. No guarda: eso lo hace quien llama.

    El índice de textos se actualiza a mano después de cada corte: rehacerlo entero
    por cada faltante costaba minutos en las tesis con cientos de referencias.
    """
    index = [list(entry) for entry in _reference_flat_index(doc)]
    done = created = 0
    for item in sorted(_open_missing(doc), key=lambda x: (int(x.get("page") or 0), float(x.get("y") or 0))):
        needle = _flatten_with_map(_norm(item.get("first_line")))[0][:MISSING_NEEDLE]
        if len(needle) < 25:
            continue
        position = next(((i, entry[1].find(needle)) for i, entry in enumerate(index) if entry[1].find(needle) > 0), None)
        if position is None:
            continue
        slot, at = position
        ref, _flat, positions = index[slot]
        text = _norm(ref.get("text"))
        head, tail = _norm(text[: positions[at]]), _norm(text[positions[at]:])
        if not head or not tail:
            continue
        # _visible_refs entrega copias: para partir hay que agarrar la referencia viva.
        parts = _apply_split(doc, _get_ref(doc, ref["id"]), [head, tail])
        index[slot:slot + 1] = [[part, *_flatten_with_map(_norm(part.get("text")))] for part in parts]
        doc.setdefault("audit_decisions", {})[item["id"]] = {
            "decision": "resolved", "note": "split_host_bulk", "timestamp": _now(),
        }
        done += 1
        created += len(parts) - 1
    if done:
        _REFERENCE_FLAT_CACHE.pop(doc.get("id"), None)
        _log_correction(doc, "split_all_merged", {"resolved": done, "references_created": created})
        doc["updated_at"] = _now()
    return done, created


DATASET_README = """# Dataset de avisos y extracción

Una fila por aviso, por tesis, por página y por referencia, para buscar patrones y
afinar el extractor y el parser. Todo sale de lo que ya está guardado: no se vuelve
a leer ningún PDF.

## tesis.csv
Una fila por tesis. `metodo_deteccion`, `estilo`, `refs_por_100_paginas`,
`ratio_tokens_corruptos` y `paginas_seccion_sin_referencias` son las columnas que
más rápido delatan una extracción rota: una tesis con páginas de bibliografía sin
ninguna referencia extraída casi siempre tiene el margen de entrada mal detectado.
`x0_moda` es el margen izquierdo más común entre sus referencias y `x0_moda_parte`
qué proporción lo comparte; si esa proporción es baja, los registros empezaron en
cualquier lado.

## secciones.csv
Una fila por sección de bibliografía detectada. `sospechosa` se marca cuando sus
referencias promedian más de 600 caracteres o cuando hay menos de 1,5 por página:
las dos formas en que aparece una zona falsa (un «referencias bibliográficas» que
estaba en medio del texto y se tragó páginas enteras de prosa).

## avisos.csv
Una fila por aviso abierto. `tipo` vale:
- `pegada_divisible`: el párrafo está adentro de una referencia ya extraída y el corte cae adentro (se divide).
- `pegada_ya_extraida`: el párrafo es el comienzo de esa referencia; no hay nada que dividir.
- `falta_de_verdad`: no está en ninguna referencia.

## paginas.csv
Una fila por página de una sección de bibliografía, con cuántas referencias y
cuántos avisos cayeron ahí.

## referencias.csv
Una fila por referencia viva: `x0` (margen izquierdo del renglón donde empieza),
`estilo`, `evidencia_borde`, `origen`, `largo` y los campos parseados. Sirve para
ver qué parsea mal y con qué forma de entrada.
"""


def _csv_bytes(header: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def _xl(value: Any) -> Any:
    """Excel refuses control characters (an unreadable text layer is full of them)."""
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _install_xlsx_sanitiser() -> None:
    """Strip control characters in every sheet BiblioRef writes.

    A single reference carrying them (OCR, or a text layer that cannot be read)
    made the whole export fail with no file and no message, which from the
    «Exportar» screen looked like the button doing nothing.
    """
    if openpyxl is None:
        return
    from openpyxl.worksheet.worksheet import Worksheet

    if getattr(Worksheet.append, "_biblioref_sanitised", False):
        return
    original = Worksheet.append

    def append(self, iterable):
        if isinstance(iterable, (list, tuple)):
            iterable = [_xl(value) for value in iterable]
        return original(self, iterable)

    append._biblioref_sanitised = True
    Worksheet.append = append


_install_xlsx_sanitiser()


REVISION_HEADER = ["Código", "Página", "N.º ref", "Estado", "Qué hacer", "Texto actual", "Corrección propuesta",
                   "Texto perdido", "Faltan datos", "Decisión", "ID de la referencia", "ID del aviso", "ID de la tesis"]


def _revision_rows(doc: dict[str, Any]) -> tuple[list[list[Any]], Counter]:
    """Every reference of a thesis in order, with the notices placed where they belong.

    A notice inside a reference is shown on that reference's row; one that is in no
    reference gets a row of its own after the last reference that starts on or before
    its page. The proposed correction is the text as it would end up.
    """
    refs = _visible_refs(doc)
    counts: Counter = Counter()
    notes: dict[str, list[tuple[str, str, str, str, str]]] = defaultdict(list)
    loose: list[tuple[int, float, dict[str, Any]]] = []
    for item in _open_missing(doc):
        inside = _missing_inside_reference(doc, item)
        if not inside:
            loose.append((int(item.get("page") or 0), float(item.get("y") or 0), item))
            continue
        preview = _missing_preview(doc, item, inside) or {}
        lost = " ".join(preview.get("leftover") or [])
        if inside.get("cut_at") is not None:
            try:
                _, parts = _split_parts_for_missing(doc, item)
                proposal = "① " + parts[0] + "\n② " + parts[1]
            except HTTPException:
                proposal = ""
            notes[inside["reference_id"]].append(("Pegada: son dos referencias", "Dividir en ①  y ②", proposal, lost, item.get("id")))
        elif lost:
            host = next((r for r in refs if r.get("id") == inside["reference_id"]), {})
            notes[inside["reference_id"]].append(("Le falta texto", "Agregar el texto perdido al final",
                                                  _norm(f"{host.get('text', '')} {lost}"), lost, item.get("id")))
    decisions = doc.get("audit_decisions") or {}
    for suspect in (doc.get("analysis") or {}).get("coverage", {}).get("merged_suspects") or []:
        if (decisions.get(f"merged_{suspect.get('page') or 0}_{suspect.get('record_index') or 0}") or {}).get("decision") in {"ignore", "resolved"}:
            continue
        head = _norm(suspect.get("text"))[:60]
        host = next((r for r in refs if _norm(r.get("text")).startswith(head)), None)
        if host:
            notes[host["id"]].append(("Posible fusionada", "Revisar si son varias referencias", "", "",
                                      f"merged_{suspect.get('page') or 0}_{suspect.get('record_index') or 0}"))
    rows: list[list[Any]] = []
    loose.sort(key=lambda x: (x[0], x[1]))
    for index, ref in enumerate(refs):
        parsed = {**(ref.get("parsed") or {}), **((ref.get("parsed_review") or {}).get("overrides") or {})}
        missing = [label for key, label in (("publication_type", "tipo"), ("author", "autor"), ("year", "año"))
                   if not str(parsed.get(key) or "").strip() or (key == "publication_type" and parsed.get(key) == "sin_clasificar")]
        found = notes.get(ref.get("id")) or []
        if found:
            for estado, hacer, proposal, lost, check in found:
                counts[estado] += 1
                rows.append([doc.get("thesis_id"), ref.get("start_page"), ref.get("order"), estado, hacer, _norm(ref.get("text")),
                             proposal, lost, ", ".join(missing), "", ref.get("id"), check, doc.get("id")])
        else:
            estado = "Faltan datos" if missing else "Sana"
            counts[estado] += 1
            rows.append([doc.get("thesis_id"), ref.get("start_page"), ref.get("order"), estado,
                         "Completar " + ", ".join(missing) if missing else "", _norm(ref.get("text")), "", "",
                         ", ".join(missing), "", ref.get("id"), "", doc.get("id")])
        next_page = int(refs[index + 1].get("start_page") or 0) if index + 1 < len(refs) else 10 ** 9
        while loose and loose[0][0] < next_page:
            page, _, item = loose.pop(0)
            counts["Falta: no se extrajo"] += 1
            rows.append([doc.get("thesis_id"), page, None, "Falta: no se extrajo", "Agregar como referencia (revisar el texto)",
                         "", _norm(item.get("text")), "", "", "", "", item.get("id"), doc.get("id")])
    for page, _, item in loose:
        counts["Falta: no se extrajo"] += 1
        rows.append([doc.get("thesis_id"), page, None, "Falta: no se extrajo", "Agregar como referencia (revisar el texto)",
                     "", _norm(item.get("text")), "", "", "", "", item.get("id"), doc.get("id")])
    return rows, counts


_YES = {"si", "sí", "s", "x", "ok", "aceptar", "aceptado", "1", "true", "dividir", "agregar", "aplicar"}
_NO = {"no", "n", "0", "false", "rechazar", "rechazado", "dejar", "dejar asi", "dejar así", "ignorar"}


def _decision_of(value: Any) -> str | None:
    text = _norm(str(value or "")).lower()
    if text in _YES:
        return "yes"
    if text in _NO:
        return "no"
    return None


def _same_text(a: Any, b: Any) -> bool:
    return _flatten_with_map(_norm(a))[0] == _flatten_with_map(_norm(b))[0]


def _close_check(doc: dict[str, Any], check_id: str, decision: str, note: str) -> None:
    doc.setdefault("audit_decisions", {})[check_id] = {"decision": decision, "note": note, "timestamp": _now()}
    _log_correction(doc, "audit_check_decision", {"check_id": check_id, "decision": decision, "note": note})


@app.post("/api/projects/{project_id}/revision/import")
async def import_revision_xlsx(project_id: str, file: UploadFile = File(...)):
    """Aplica las decisiones anotadas en la columna «Decisión» del Excel de revisión.

    «sí» aplica la corrección propuesta (tal como quedó escrita, si se la editó) y «no»
    cierra el aviso. Una fila cuya referencia cambió desde la exportación no se aplica:
    el texto de la corrección ya no correspondería a lo que hay.
    """
    if openpyxl is None:
        raise HTTPException(500, "Falta openpyxl para leer Excel")
    _ensure_projects()
    _get_project(project_id)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
        ws = wb["Revisión"]
    except Exception:
        raise HTTPException(400, "No es un Excel de revisión de BiblioRef (falta la hoja «Revisión»)")
    rows = ws.iter_rows(values_only=True)
    header = [str(v or "").strip() for v in next(rows, [])]
    missing_cols = [name for name in REVISION_HEADER if name not in header]
    if missing_cols:
        raise HTTPException(400, "Faltan columnas: " + ", ".join(missing_cols))
    col = {name: header.index(name) for name in REVISION_HEADER}
    applied: Counter = Counter()
    skipped: list[dict[str, Any]] = []
    touched: set[str] = set()
    last_ref: dict[str, str] = {}
    replaced: dict[str, str] = {}  # a split reference lives on in its second part

    def skip(row_number: int, code: Any, why: str) -> None:
        skipped.append({"fila": row_number, "tesis": code, "motivo": why})

    for row_number, row in enumerate(rows, start=2):
        get = lambda name: row[col[name]] if col[name] < len(row) else None
        doc_id, ref_id, check_id = get("ID de la tesis"), get("ID de la referencia"), get("ID del aviso")
        estado, decision = str(get("Estado") or ""), _decision_of(get("Decisión"))
        if doc_id and ref_id:
            last_ref[doc_id] = ref_id
        if not decision:
            continue
        doc = STORE.docs().get(doc_id or "")
        if not doc or doc.get("project_id") != project_id:
            skip(row_number, get("Código"), "la tesis no está en este proyecto")
            continue
        refs = {r.get("id"): r for r in doc.get("references") or [] if not r.get("deleted")}
        open_checks = {x.get("id") for x in _open_missing(doc)}
        try:
            if decision == "no":
                decided = doc.get("audit_decisions") or {}
                ids = ([f"data_{kind}_{ref_id}" for kind in ("type", "author", "year")]
                       if estado == "Faltan datos" and ref_id in refs else [check_id] if check_id else [])
                ids = [x for x in ids if (decided.get(x) or {}).get("decision") not in {"ignore", "resolved"}]
                if not ids:
                    continue  # already closed: importing the same file twice changes nothing
                for x in ids:
                    _close_check(doc, x, "ignore", "revision_excel")
                applied["cerrados sin cambios"] += 1
            elif estado == "Pegada: son dos referencias":
                ref = refs.get(ref_id)
                if not ref or check_id not in open_checks:
                    skip(row_number, doc.get("thesis_id"), "el aviso ya estaba resuelto")
                    continue
                if not _same_text(ref.get("text"), get("Texto actual")):
                    skip(row_number, doc.get("thesis_id"), f"REF {ref.get('order')} cambió desde la exportación")
                    continue
                proposal = str(get("Corrección propuesta") or "")
                parts = [_norm(x) for x in re.split(r"[①②]", proposal) if _norm(x)]
                if len(parts) != 2:
                    skip(row_number, doc.get("thesis_id"), "la corrección no tiene las dos partes ① y ②")
                    continue
                new_refs = _apply_split(doc, ref, parts)
                replaced[ref_id] = new_refs[-1]["id"]
                _close_check(doc, check_id, "resolved", "split_from_revision_excel")
                applied["divididas"] += 1
            elif estado == "Le falta texto":
                ref = refs.get(ref_id)
                if not ref or check_id not in open_checks:
                    skip(row_number, doc.get("thesis_id"), "el aviso ya estaba resuelto")
                    continue
                if not _same_text(ref.get("text"), get("Texto actual")):
                    skip(row_number, doc.get("thesis_id"), f"REF {ref.get('order')} cambió desde la exportación")
                    continue
                before = ref["text"]
                ref["text"] = _norm(get("Corrección propuesta"))
                ref.pop("parsed_review", None)
                _reset_reference_resolution(ref, preserve_human=False)
                _refresh_reference_parse(ref)
                ref["edited"] = True
                _log_correction(doc, "edit_reference", {"reference_id": ref["id"], "before": before,
                                                        "after": ref["text"], "source": "revision_excel"})
                _close_check(doc, check_id, "resolved", "appended_from_revision_excel")
                applied["completadas"] += 1
            elif estado == "Falta: no se extrajo":
                text = _norm(get("Corrección propuesta"))
                if check_id not in open_checks or not text:
                    skip(row_number, doc.get("thesis_id"), "el aviso ya estaba resuelto o el texto está vacío")
                    continue
                head = _flatten_with_map(text)[0][:120]
                twin = next((r for r in refs.values() if len(head) >= 20 and difflib.SequenceMatcher(
                    None, head, _flatten_with_map(_norm(r.get("text")))[0][:120]).ratio() >= 0.85), None)
                if twin:
                    skip(row_number, doc.get("thesis_id"), f"ya existe como REF {twin.get('order')}: se cerró el aviso sin duplicarla")
                    _close_check(doc, check_id, "resolved", "already_present_revision_excel")
                    touched.add(doc_id)
                    continue
                ref = {
                    "id": f"ref_{uuid.uuid4().hex[:12]}", "order": 0, "text": text, "original_text": "",
                    "start_page": int(get("Página") or 0) or None, "end_page": int(get("Página") or 0) or None,
                    "start_y": None, "end_y": None, "start_x": None, "zone_id": None, "style": None,
                    "boundary_evidence": "", "audit_flags": "", "status": "reviewed", "source": "human_insert",
                    "edited": True, "deleted": False,
                }
                _refresh_reference_parse(ref)
                seq = doc.setdefault("references", [])
                anchor_id = last_ref.get(doc_id)
                while anchor_id in replaced:
                    anchor_id = replaced[anchor_id]
                after = next((r for r in seq if r.get("id") == anchor_id), None)
                seq.insert(seq.index(after) + 1 if after in seq else len(seq), ref)
                _log_correction(doc, "insert_reference", {"reference_id": ref["id"], "text": text,
                                                          "page": ref["start_page"], "source": "revision_excel"})
                _close_check(doc, check_id, "resolved", "reference_added_from_revision_excel")
                last_ref[doc_id] = ref["id"]
                applied["agregadas"] += 1
            elif estado == "Posible fusionada":
                skip(row_number, doc.get("thesis_id"), "una posible fusionada se divide a mano en Referencias por tesis")
                continue
            else:
                continue
            touched.add(doc_id)
        except HTTPException as exc:
            skip(row_number, doc.get("thesis_id"), str(exc.detail))
    for doc_id in touched:
        doc = STORE.docs()[doc_id]
        _renumber(doc)
        _REFERENCE_FLAT_CACHE.pop(doc_id, None)
        doc["updated_at"] = _now()
    if touched:
        STORE.save()
    return {"applied": dict(applied), "n_applied": sum(applied.values()), "skipped": skipped,
            "documents": sorted(STORE.docs()[d].get("thesis_id") or d for d in touched)}


@app.get("/api/projects/{project_id}/revision.xlsx")
def project_revision_xlsx(project_id: str, document_id: str | None = None):
    """Todas las referencias con sus avisos intercalados, para revisar en Excel sin recortes."""
    if openpyxl is None:
        raise HTTPException(500, "Falta openpyxl para exportar a Excel")
    from openpyxl.styles import Alignment, Font, PatternFill
    _ensure_projects()
    project = _get_project(project_id)
    docs = [d for d in STORE.docs().values() if d.get("project_id") == project_id and not d.get("deleted")
            and (not document_id or d.get("id") == document_id)]
    docs.sort(key=lambda d: str(d.get("thesis_id") or ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "Revisión"
    ws.append(REVISION_HEADER)
    summary = wb.create_sheet("Resumen")
    kinds = ["Sana", "Faltan datos", "Pegada: son dos referencias", "Le falta texto", "Posible fusionada", "Falta: no se extrajo"]
    summary.append(["Código", "Título de la tesis", "Autor/a", "Año", "Referencias"] + kinds)
    amber = PatternFill("solid", fgColor="FDF1E2")
    # The row number is carried along: asking openpyxl for ws.max_row on every row
    # makes it scan the whole sheet again (minutes on a project of 120 theses).
    at = 1
    for doc in docs:
        rows, counts = _revision_rows(doc)
        for row in rows:
            at += 1
            ws.append([_xl(v) for v in row])
            if row[3] != "Sana":
                for column in range(1, len(REVISION_HEADER) + 1):
                    ws.cell(row=at, column=column).fill = amber
        metadata = doc.get("metadata") or {}
        summary.append([_xl(v) for v in [doc.get("thesis_id"), _doc_title(doc), metadata.get("author"), metadata.get("year"),
                        len(_visible_refs(doc))] + [counts.get(k, 0) for k in kinds]])
    from openpyxl.utils import get_column_letter
    for sheet, widths, last_row in ((ws, (8, 7, 7, 24, 30, 70, 70, 28, 14, 14, 18, 20, 18), at),
                                    (summary, (8, 50, 24, 7, 11, 8, 12, 14, 12, 12, 14), len(docs) + 1)):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{max(last_row, 1)}"
        for column, width in zip("ABCDEFGHIJKLM", widths):
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    wb.save(buffer)
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(project.get("name") or project_id)).strip("_") or "proyecto"
    return Response(content=buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="revision_{name}.xlsx"', "Cache-Control": "no-store"})


@app.get("/api/projects/{project_id}/avisos.xlsx")
def project_notices_xlsx(project_id: str):
    """Faltantes y pegadas de todas las tesis, en un Excel para revisar o compartir."""
    if openpyxl is None:
        raise HTTPException(500, "Falta openpyxl para exportar a Excel")
    _ensure_projects()
    project = _get_project(project_id)
    docs = [d for d in STORE.docs().values() if d.get("project_id") == project_id and not d.get("deleted")]
    docs.sort(key=lambda d: str(d.get("thesis_id") or ""))
    header = ["Código", "Título de la tesis", "Autor/a", "Año", "Página", "Tipo", "Primera línea del recuadro",
              "Texto del recuadro", "REF que lo contiene", "Texto de esa REF", "Texto perdido", "Parecido con la más cercana",
              "ID del aviso", "ID de la tesis"]
    missing_rows: list[list[Any]] = []
    glued_rows: list[list[Any]] = []
    kind_label = {"falta_de_verdad": "Falta en las referencias",
                  "pegada_divisible": "Pegada a otra referencia (se puede dividir)",
                  "pegada_ya_extraida": "Continuación o ya extraída"}
    for doc in docs:
        metadata = doc.get("metadata") or {}
        refs = {r.get("id"): r for r in _visible_refs(doc)}
        for item in _open_missing(doc):
            inside = _missing_inside_reference(doc, item)
            kind = ("pegada_ya_extraida" if inside and inside.get("cut_at") is None
                    else "pegada_divisible" if inside else "falta_de_verdad")
            preview = _missing_preview(doc, item, inside) if inside else None
            host = refs.get((inside or {}).get("reference_id")) or {}
            row = [
                doc.get("thesis_id"), _doc_title(doc), metadata.get("author"), metadata.get("year"),
                int(item.get("page") or 0), kind_label[kind], _norm(item.get("first_line")), _norm(item.get("text")),
                (inside or {}).get("order"), _norm(host.get("text")) if host else "",
                " … ".join((preview or {}).get("leftover") or []), item.get("best_existing_match"),
                item.get("id"), doc.get("id"),
            ]
            (glued_rows if inside else missing_rows).append(row)
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in (("Faltantes", missing_rows), ("Pegadas", glued_rows)):
        ws = wb.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append([_xl(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column, width in zip("ABCDEFGHIJKLMN", (9, 40, 22, 7, 8, 30, 50, 70, 10, 70, 30, 12, 20, 18)):
            ws.column_dimensions[column].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(project.get("name") or project_id)).strip("_") or "proyecto"
    return Response(content=buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="avisos_{name}.xlsx"', "Cache-Control": "no-store"})


@app.get("/api/projects/{project_id}/dataset.zip")
def project_dataset(project_id: str, include_references: bool = True):
    """Lo que sabe BiblioRef de sus propios errores, en CSV, para refinar extractor y parser."""
    _ensure_projects()
    project = _get_project(project_id)
    docs = [d for d in STORE.docs().values() if d.get("project_id") == project_id and not d.get("deleted")]
    docs.sort(key=lambda d: str(d.get("thesis_id") or ""))

    tesis_rows: list[list[Any]] = []
    secciones_rows: list[list[Any]] = []
    avisos_rows: list[list[Any]] = []
    paginas_rows: list[list[Any]] = []
    refs_rows: list[list[Any]] = []

    for doc in docs:
        analysis = doc.get("analysis") or {}
        coverage = analysis.get("coverage") or {}
        plausibility = analysis.get("plausibility") or {}
        metadata = doc.get("metadata") or {}
        decisions = doc.get("audit_decisions") or {}
        refs = _visible_refs(doc)
        sections = doc.get("sections") or []
        section_pages = set()
        for section in sections:
            first = int(section.get("start_page") or 0)
            last = int(section.get("end_page") or first)
            section_pages.update(range(first, last + 1))

        refs_by_page = Counter(int(r.get("start_page") or 0) for r in refs)
        x0_values = [round(float(r["start_x"]), 1) for r in refs if r.get("start_x") is not None]
        x0_mode, x0_share = (None, None)
        if x0_values:
            x0_mode, count = Counter(x0_values).most_common(1)[0]
            x0_share = round(count / len(x0_values), 3)

        kinds = Counter()
        page_avisos = Counter()
        for item in _open_missing(doc):
            inside = _missing_inside_reference(doc, item)
            kind = ("pegada_ya_extraida" if inside and inside.get("cut_at") is None
                    else "pegada_divisible" if inside else "falta_de_verdad")
            kinds[kind] += 1
            page_avisos[(int(item.get("page") or 0), kind)] += 1
            avisos_rows.append([
                doc.get("thesis_id"), doc.get("id"), item.get("id"), item.get("page"), item.get("y"), kind,
                _norm(item.get("first_line")), _norm(item.get("text")), item.get("best_existing_match"),
                (inside or {}).get("order"), (inside or {}).get("length"), (inside or {}).get("cut_at"),
                (decisions.get(item.get("id")) or {}).get("decision") or "",
            ])

        empty_section_pages = sum(1 for page in section_pages if not refs_by_page.get(page))
        tesis_rows.append([
            doc.get("thesis_id"), doc.get("id"), _doc_title(doc), metadata.get("author"), metadata.get("year"),
            doc.get("page_count"), len(refs), plausibility.get("references_per_100_pages"),
            plausibility.get("corrupted_token_ratio"), plausibility.get("implausible"),
            analysis.get("detection_method"), analysis.get("algorithm_version"), analysis.get("software_version"),
            analysis.get("extraction_revision"), analysis.get("status"), bool(analysis.get("requires_confirmation")),
            "; ".join(f"{sec.get('start_page')}-{sec.get('end_page')}" for sec in sections),
            len(section_pages), empty_section_pages, x0_mode, x0_share,
            (coverage.get("summary") or {}).get("n_possible_missing"),
            (coverage.get("summary") or {}).get("n_merged_suspects"),
            (coverage.get("summary") or {}).get("n_page_boundaries_review"),
            kinds["pegada_divisible"], kinds["pegada_ya_extraida"], kinds["falta_de_verdad"],
        ])

        # Una sección cuyas referencias pesan miles de caracteres no es una bibliografía:
        # casi siempre es un encabezado encontrado en el cuerpo del texto.
        for section in sections:
            first = int(section.get("start_page") or 0)
            last = int(section.get("end_page") or first)
            inside = [r for r in refs if first <= int(r.get("start_page") or 0) <= last]
            lengths = [len(_norm(r.get("text"))) for r in inside]
            average = round(sum(lengths) / len(lengths)) if lengths else 0
            per_page = round(len(inside) / max(1, last - first + 1), 2)
            secciones_rows.append([
                doc.get("thesis_id"), section.get("label"), section.get("zone_id"), first, last,
                last - first + 1, len(inside), per_page, average, max(lengths) if lengths else 0,
                "sí" if (average > 600 or per_page < 1.5) else "no",
            ])

        for page in sorted(section_pages):
            paginas_rows.append([
                doc.get("thesis_id"), page, refs_by_page.get(page, 0),
                page_avisos.get((page, "pegada_divisible"), 0) + page_avisos.get((page, "pegada_ya_extraida"), 0),
                page_avisos.get((page, "falta_de_verdad"), 0),
            ])

        if include_references:
            for ref in refs:
                parsed = _audit_parsed(ref)
                refs_rows.append([
                    doc.get("thesis_id"), ref.get("id"), ref.get("order"), ref.get("start_page"), ref.get("end_page"),
                    ref.get("start_x"), ref.get("style"), ref.get("boundary_evidence"), ref.get("source"),
                    ref.get("status"), len(_norm(ref.get("text"))), _audit_scalar(parsed.get("author")),
                    _audit_scalar(parsed.get("year")), _audit_scalar(parsed.get("publication_type")),
                    "sí" if _norm(parsed.get("doi")) else "no", _norm(ref.get("text")),
                ])

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("README.md", DATASET_README)
        bundle.writestr("tesis.csv", _csv_bytes([
            "tesis_id", "document_id", "titulo", "autor", "anio", "paginas", "n_referencias",
            "refs_por_100_paginas", "ratio_tokens_corruptos", "implausible", "metodo_deteccion",
            "algoritmo", "software", "revision_extraccion", "estado", "requiere_confirmacion",
            "secciones", "paginas_seccion", "paginas_seccion_sin_referencias", "x0_moda", "x0_moda_parte",
            "n_posibles_faltantes", "n_fusionadas_detectadas", "n_bordes_a_revisar",
            "n_pegadas_divisibles", "n_pegadas_ya_extraidas", "n_faltan_de_verdad",
        ], tesis_rows))
        bundle.writestr("secciones.csv", _csv_bytes([
            "tesis_id", "etiqueta", "zone_id", "pagina_inicio", "pagina_fin", "paginas",
            "n_referencias", "referencias_por_pagina", "largo_promedio", "largo_maximo", "sospechosa",
        ], secciones_rows))
        bundle.writestr("avisos.csv", _csv_bytes([
            "tesis_id", "document_id", "aviso_id", "pagina", "y", "tipo", "primera_linea", "texto",
            "parecido_mas_cercano", "referencia_orden", "referencia_largo", "corte_en", "decision_humana",
        ], avisos_rows))
        bundle.writestr("paginas.csv", _csv_bytes([
            "tesis_id", "pagina", "n_referencias", "n_pegadas", "n_faltan_de_verdad",
        ], paginas_rows))
        if include_references:
            bundle.writestr("referencias.csv", _csv_bytes([
                "tesis_id", "referencia_id", "orden", "pagina_inicio", "pagina_fin", "x0", "estilo",
                "evidencia_borde", "origen", "estado", "largo", "autor", "anio", "tipo", "tiene_doi", "texto",
            ], refs_rows))
    name = _safe_name(str(project.get("name") or project_id))
    return Response(
        content=buffer.getvalue(), media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="dataset_{name}.zip"', "Cache-Control": "no-store"},
    )


@app.post("/api/projects/{project_id}/missing/split-all")
def split_all_merged(project_id: str, document_id: str | None = None):
    """El «Dividir todas» del Paso 1: un solo guardado para todo el proyecto."""
    _ensure_projects()
    _get_project(project_id)
    docs = [d for d in STORE.docs().values()
            if d.get("project_id") == project_id and not d.get("deleted")
            and (not document_id or d.get("id") == document_id)]
    resolved = created = touched = 0
    for doc in docs:
        done, made = _split_all_in_document(doc)
        if done:
            touched += 1
            resolved += done
            created += made
    if resolved:
        STORE.save()
    return {"ok": True, "resolved": resolved, "references_created": created, "documents": touched}


@app.get("/api/projects/{project_id}/missing")
def project_missing(
    project_id: str,
    document_id: str | None = None,
    q: str = "",
    order: str = "thesis",
    stage: str = "all",
    offset: int = 0,
    limit: int = 60,
):
    """Every open possible missing reference of a project, grouped by thesis."""
    needle = _norm(q).lower()
    limit = min(max(int(limit), 1), 200)
    offset = max(int(offset), 0)
    _ensure_projects()
    _get_project(project_id)
    docs = [d for d in STORE.docs().values()
            if d.get("project_id") == project_id and not d.get("deleted")]
    docs.sort(key=lambda d: str(d.get("thesis_id") or ""))
    flat: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any] | None]] = []
    per_doc: dict[str, int] = {}
    # Fusions read from the text (the sidebar's «posibles fusionadas») are a different
    # signal from paragraphs found inside a reference; the screen reports both.
    n_merged_suspects = 0
    for doc in docs:
        if document_id and doc.get("id") != document_id:
            continue
        n_merged_suspects += problem_counts(doc)["merged"]
        items = _open_missing(doc)
        per_doc[doc["id"]] = len(items)
        for item in items:
            if needle:
                hay = " ".join([_norm(item.get("text")), _norm(item.get("first_line")),
                                str(doc.get("thesis_id") or ""), _norm(_doc_title(doc))]).lower()
                if needle not in hay:
                    continue
            flat.append((doc, item, _missing_inside_reference(doc, item)))
    total_all = len(flat)
    n_inside = sum(1 for row in flat if row[2])
    # Las que están adentro pero empiezan donde empieza la referencia no se dividen:
    # el párrafo ya está extraído y lo único que queda es cerrar el aviso.
    n_splittable = sum(1 for row in flat if row[2] and row[2].get("cut_at") is not None)
    # Extraer un párrafo que ya está adentro de otra referencia lo duplicaría: primero se divide.
    if stage == "merged":
        flat = [row for row in flat if row[2]]
    elif stage == "real":
        flat = [row for row in flat if not row[2]]
    if order == "likely":
        flat.sort(key=lambda row: (float(row[1].get("best_existing_match") or 0),
                                   str(row[0].get("thesis_id") or ""), int(row[1].get("page") or 0)))
    else:
        flat.sort(key=lambda row: (str(row[0].get("thesis_id") or ""),
                                   int(row[1].get("page") or 0), float(row[1].get("y") or 0)))
    page_items = flat[offset:offset + limit]
    groups: dict[str, dict[str, Any]] = {}
    for doc, item, inside in page_items:
        group = groups.get(doc["id"])
        if not group:
            metadata = doc.get("metadata") or {}
            section_pages: list[int] = []
            for sec in doc.get("sections") or []:
                first_page = max(1, int(sec.get("start_page") or 1))
                last_page = min(int(doc.get("page_count") or first_page), int(sec.get("end_page") or first_page))
                section_pages.extend(range(first_page, last_page + 1))
            refs_by_page = Counter(int(r.get("start_page") or 0) for r in _visible_refs(doc))
            avisos_by_page: Counter = Counter()
            merged_by_page: Counter = Counter()
            for open_item in _open_missing(doc):
                page_number = int(open_item.get("page") or 0)
                avisos_by_page[page_number] += 1
                if _missing_inside_reference(doc, open_item):
                    merged_by_page[page_number] += 1
            group = groups[doc["id"]] = {
                "document_id": doc["id"],
                "pages": [{"page": page_number,
                           "references": refs_by_page.get(page_number, 0),
                           "merged": merged_by_page.get(page_number, 0),
                           "real": avisos_by_page.get(page_number, 0) - merged_by_page.get(page_number, 0)}
                          for page_number in sorted(dict.fromkeys(section_pages))],
                "thesis_id": doc.get("thesis_id"),
                "title": _doc_title(doc),
                "author": metadata.get("author") or "",
                "year": metadata.get("year") or "",
                "page_count": doc.get("page_count"),
                "n_references": doc.get("n_references"),
                "n_missing": per_doc.get(doc["id"], 0),
                "sections": [{k: sec.get(k) for k in ("label", "start_page", "end_page", "start_y", "end_y", "zone_id")}
                             for sec in doc.get("sections") or []],
                "items": [],
            }
        group["items"].append({
            "id": item.get("id"),
            "page": int(item.get("page") or 0),
            "y": float(item.get("y") or 0),
            "first_line": _norm(item.get("first_line")),
            "text": _norm(item.get("text"))[:320],
            "best_existing_match": item.get("best_existing_match"),
            "inside_ref": inside,
            "preview": _missing_preview(doc, item, inside),
        })
    return {
        "project_id": project_id,
        "total": len(flat),
        "offset": offset,
        "returned": len(page_items),
        "stage": stage,
        "n_inside_reference": n_inside,
        "n_splittable": n_splittable,
        "n_real": total_all - n_inside,
        "n_merged_suspects": n_merged_suspects,
        "n_documents_with_missing": sum(1 for n in per_doc.values() if n),
        "documents": list(groups.values()),
    }


@app.get("/api/documents/{doc_id}/style-profile")
def style_profile_document(doc_id: str, include_references: bool = True):
    """Shadow-only style morphology profile; never mutates extraction/parser state."""
    doc = STORE.get(doc_id)
    refs = [_reference_with_parse(r) for r in doc.get("references", []) if not r.get("deleted")]
    context = {
        "languages": None,
        "country": (doc.get("metadata") or {}).get("country"),
        "institution": (doc.get("metadata") or {}).get("university"),
        "discipline": (doc.get("metadata") or {}).get("area"),
        "publication_year": (doc.get("metadata") or {}).get("year"),
    }
    return profile_style_document(refs, context=context, include_references=include_references)


@app.get("/api/documents/{doc_id}/style-profile.json")
def export_style_profile(doc_id: str):
    """Downloadable JSON sidecar for the shadow StyleProfiler layer."""
    profile = style_profile_document(doc_id, include_references=True)
    payload = json.dumps(profile, ensure_ascii=False, indent=2).encode("utf-8")
    return Response(
        content=payload,
        media_type="application/json; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="biblioref_style_profile_{doc_id}.json"'},
    )


@app.get("/api/documents/{doc_id}/export/{fmt}")
def export_document(doc_id: str, fmt: str):
    doc = STORE.get(doc_id)
    refs = _visible_refs(doc)
    stem = Path(doc["filename"]).stem
    if fmt == "json":
        payload = {
            "document": _public_doc(doc, False),
            "references": refs,
            "sections": doc.get("sections", []),
            "corrections": doc.get("corrections", []),
            "metadata_inference": doc.get("metadata_inference", {}),
        }
        return Response(
            json.dumps(payload, ensure_ascii=False, indent=2), media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{stem}_biblioref.json"'},
        )
    if fmt == "csv":
        out = io.StringIO()
        fieldnames = ["thesis_id", "thesis_title", "filename", "document_id", "reference_id", "author", "year", "university", "faculty", "program", "repository", "area", "country", "metadata_inferred", "order", "text", "start_page", "end_page", "zone_id", "status", "source", "edited", *PARSED_EXPORT_FIELDS]
        writer = csv.DictWriter(out, fieldnames=fieldnames)
        writer.writeheader()
        for r in refs:
            writer.writerow({
                "thesis_id": doc.get("thesis_id"), "thesis_title": doc.get("title"),
                "author": (doc.get("metadata") or {}).get("author"), "year": (doc.get("metadata") or {}).get("year"),
                "university": (doc.get("metadata") or {}).get("university"), "faculty": (doc.get("metadata") or {}).get("faculty"),
                "program": (doc.get("metadata") or {}).get("program"), "repository": (doc.get("metadata") or {}).get("repository"),
                "area": (doc.get("metadata") or {}).get("area"), "country": (doc.get("metadata") or {}).get("country"),
                "filename": doc.get("filename"), "document_id": doc.get("id"), "reference_id": r.get("id"),
                "metadata_inferred": _inferred_fields(doc),
                **{k: r.get(k) for k in ["order", "text", "start_page", "end_page", "zone_id", "status", "source", "edited"]},
                **_flat_parse_export(r),
            })
        data = out.getvalue().encode("utf-8-sig")
        return Response(data, media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{stem}_biblioref.csv"'})
    if fmt == "xlsx":
        if Workbook is None:
            raise HTTPException(501, "openpyxl no está instalado")
        wb = Workbook()
        ws = wb.active
        ws.title = "Referencias"
        headers = ["Tesis ID", "Título de tesis", "Archivo", "Referencia ID", "N", "Referencia", "Página inicial", "Página final", "Sección", "Estado", "Fuente de extracción", "Editada", "Autor parseado", "Autores parseados", "Año parseado", "Tipo documental", "Título parseado", "Revista", "Editorial normalizada", "Editorial observada", "DOI", "Parser"]
        ws.append(headers)
        for r in refs:
            pmeta = (r.get("parsed") or {})
            ws.append([doc.get("thesis_id"), doc.get("title"), doc.get("filename"), r.get("id"), r.get("order"), r.get("text"), r.get("start_page"), r.get("end_page"), r.get("zone_id"), r.get("status"), r.get("source"), bool(r.get("edited")), pmeta.get("author"), pmeta.get("authors"), pmeta.get("year"), pmeta.get("publication_type"), pmeta.get("title"), pmeta.get("journal"), pmeta.get("publisher"), pmeta.get("publisher_observed"), pmeta.get("doi"), pmeta.get("parser_version")])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 55
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 8
        ws.column_dimensions["F"].width = 100
        for col in "GHIJKL": ws.column_dimensions[col].width = 18
        sec = wb.create_sheet("Secciones")
        sec.append(["Sección", "Inicio pág.", "Inicio Y", "Fin pág.", "Fin Y", "Tipo", "Fuente", "Confirmada"])
        for s in doc.get("sections", []):
            sec.append([s.get("label"), s.get("start_page"), s.get("start_y"), s.get("end_page"), s.get("end_y"), s.get("role"), s.get("source"), s.get("confirmed")])
        corr = wb.create_sheet("Revisión")
        corr.append(["Tipo", "Fecha", "Detalle JSON"])
        for c in doc.get("corrections", []):
            corr.append([c.get("kind"), c.get("timestamp"), json.dumps(c, ensure_ascii=False)])
        res = wb.create_sheet("Resumen")
        meta = doc.get("metadata") or {}
        res.append(["Tesis ID", doc.get("thesis_id")]); res.append(["Título", doc.get("title")]); res.append(["Autor/a", meta.get("author")]); res.append(["Año", meta.get("year")]); res.append(["Universidad", meta.get("university")]); res.append(["Facultad / unidad académica", meta.get("faculty")]); res.append(["Programa / carrera", meta.get("program")]); res.append(["Repositorio", meta.get("repository")]); res.append(["URL del registro", meta.get("repository_url")]); res.append(["Área / disciplina", meta.get("area")]); res.append(["País", meta.get("country")]); res.append(["Archivo", doc["filename"]]); res.append(["Páginas", doc["page_count"]]); res.append(["Referencias", len(refs)])
        res.append(["BiblioRef", SOFTWARE_VERSION]); res.append(["Algoritmo", ALGORITHM_VERSION]); res.append(["Motor estable", STABLE_ENGINE_VERSION]); res.append(["Parser bibliográfico", PARSER_VERSION]); res.append(["Estado del parser", PARSER_STATUS])
        buff = io.BytesIO(); wb.save(buff); buff.seek(0)
        return StreamingResponse(buff, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{stem}_biblioref.xlsx"'})
    raise HTTPException(404, "Formato no disponible")


def _project_docs(documents: str | None = None) -> list[dict[str, Any]]:
    """Las tesis a exportar, con su código y su título.

    Sólo se completa lo que falta, y sólo en las tesis elegidas: ``_ensure_doc_identity``
    vuelve a leer la portada del PDF cuando la tesis quedó con una revisión vieja, así
    que hacerlo sobre todo el workspace dejaba la exportación minutos sin responder
    (el relector de portadas ya hace ese trabajo en segundo plano).
    """
    ordered = sorted(STORE.docs().values(), key=lambda d: d.get("created_at", 0))
    if documents:
        wanted = {x.strip() for x in documents.split(",") if x.strip()}
        ordered = [d for d in ordered if d.get("id") in wanted]
    incomplete = [d for d in ordered if not d.get("thesis_id") or not d.get("title")]
    if any([_ensure_doc_identity(d) for d in incomplete]):
        STORE.save()
    if not ordered:
        raise HTTPException(400, "No hay tesis seleccionadas para exportar")
    return ordered


@app.get("/api/project/export/selection/{fmt}")
def export_project_selection(fmt: str, documents: str | None = None):
    docs = _project_docs(documents)
    n_docs = len(docs)
    n_refs = sum(len(_visible_refs(d)) for d in docs)
    filename_base = f"biblioref_{n_docs}_tesis"

    if fmt == "csv":
        out = io.StringIO()
        fieldnames = ["thesis_id", "thesis_title", "filename", "document_id", "reference_id", "author", "year", "university", "faculty", "program", "repository", "area", "country", "metadata_inferred", "order", "text", "start_page", "end_page", "zone_id", "status", "source", "edited", *PARSED_EXPORT_FIELDS]
        writer = csv.DictWriter(out, fieldnames=fieldnames)
        writer.writeheader()
        for doc in docs:
            for r in _visible_refs(doc):
                writer.writerow({
                    "thesis_id": doc.get("thesis_id"), "thesis_title": doc.get("title"),
                    "author": (doc.get("metadata") or {}).get("author"), "year": (doc.get("metadata") or {}).get("year"),
                    "university": (doc.get("metadata") or {}).get("university"), "faculty": (doc.get("metadata") or {}).get("faculty"),
                    "program": (doc.get("metadata") or {}).get("program"), "repository": (doc.get("metadata") or {}).get("repository"),
                    "area": (doc.get("metadata") or {}).get("area"), "country": (doc.get("metadata") or {}).get("country"),
                    "filename": doc.get("filename"), "document_id": doc.get("id"), "reference_id": r.get("id"),
                    "metadata_inferred": _inferred_fields(doc),
                    **{k: r.get(k) for k in ["order", "text", "start_page", "end_page", "zone_id", "status", "source", "edited"]},
                    **_flat_parse_export(r),
                })
        return Response(out.getvalue().encode("utf-8-sig"), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'})

    if fmt == "json":
        payload = {
            "schema": 2,
            "project": {"n_documents": n_docs, "n_references": n_refs, "software_version": SOFTWARE_VERSION},
            "documents": [
                {
                    "document": _public_doc(doc, False),
                    "references": _visible_refs(doc),
                    "sections": doc.get("sections", []),
                    "corrections": doc.get("corrections", []),
            "metadata_inference": doc.get("metadata_inference", {}),
                } for doc in docs
            ],
        }
        return Response(json.dumps(payload, ensure_ascii=False, indent=2), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{filename_base}.json"'})

    if fmt == "xlsx":
        if Workbook is None:
            raise HTTPException(501, "openpyxl no está instalado")
        wb = Workbook()
        ws = wb.active
        ws.title = "Referencias"
        ws.append(["Tesis ID", "Título de tesis", "Archivo", "Documento ID", "Referencia ID", "N", "Referencia", "Página inicial", "Página final", "Sección", "Estado", "Fuente de extracción", "Editada", "Autor parseado", "Autores parseados", "Año parseado", "Tipo documental", "Título parseado", "Revista", "Editorial normalizada", "Editorial observada", "DOI", "Parser"])
        for doc in docs:
            for r in _visible_refs(doc):
                pmeta = (r.get("parsed") or {})
                ws.append([doc.get("thesis_id"), doc.get("title"), doc.get("filename"), doc.get("id"), r.get("id"), r.get("order"), r.get("text"), r.get("start_page"), r.get("end_page"), r.get("zone_id"), r.get("status"), r.get("source"), bool(r.get("edited")), pmeta.get("author"), pmeta.get("authors"), pmeta.get("year"), pmeta.get("publication_type"), pmeta.get("title"), pmeta.get("journal"), pmeta.get("publisher"), pmeta.get("publisher_observed"), pmeta.get("doi"), pmeta.get("parser_version")])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {"A": 12, "B": 55, "C": 35, "D": 22, "E": 22, "F": 8, "G": 100, "H": 16, "I": 16, "J": 14, "K": 16, "L": 18, "M": 12}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        tws = wb.create_sheet("Tesis")
        tws.append(["Tesis ID", "Título", "Autor/a", "Año", "Universidad", "Facultad / unidad académica", "Programa / carrera", "Repositorio", "URL del registro", "Área / disciplina", "País", "Archivo", "Documento ID", "Páginas", "Referencias", "Estado", "Fuente del título", "Título confirmado", "Origen autor/a", "Origen año", "Origen universidad", "Origen facultad", "Origen programa"])
        for doc in docs:
            meta = doc.get("metadata") or {}
            tws.append([doc.get("thesis_id"), doc.get("title"), meta.get("author"), meta.get("year"), meta.get("university"), meta.get("faculty"), meta.get("program"), meta.get("repository"), meta.get("repository_url"), meta.get("area"), meta.get("country"), doc.get("filename"), doc.get("id"), doc.get("page_count"), len(_visible_refs(doc)), doc.get("status"), doc.get("title_source"), bool(doc.get("title_confirmed")), *[field_origin(doc, f) for f in ("author", "year", "university", "faculty", "program")]])

        sws = wb.create_sheet("Secciones")
        sws.append(["Tesis ID", "Título", "Sección", "Inicio pág.", "Inicio Y", "Fin pág.", "Fin Y", "Tipo", "Fuente", "Confirmada"])
        for doc in docs:
            for sec in doc.get("sections", []):
                sws.append([doc.get("thesis_id"), doc.get("title"), sec.get("label"), sec.get("start_page"), sec.get("start_y"), sec.get("end_page"), sec.get("end_y"), sec.get("role"), sec.get("source"), sec.get("confirmed")])

        cws = wb.create_sheet("Revisión")
        cws.append(["Tesis ID", "Título", "Tipo", "Fecha", "Detalle JSON"])
        for doc in docs:
            for c in doc.get("corrections", []):
                cws.append([doc.get("thesis_id"), doc.get("title"), c.get("kind"), c.get("timestamp"), json.dumps(c, ensure_ascii=False)])

        rws = wb.create_sheet("Resumen")
        rws.append(["Tesis seleccionadas", n_docs]); rws.append(["Referencias exportadas", n_refs]); rws.append(["BiblioRef", SOFTWARE_VERSION]); rws.append(["Algoritmo", ALGORITHM_VERSION]); rws.append(["Motor estable", STABLE_ENGINE_VERSION]); rws.append(["Parser bibliográfico", PARSER_VERSION]); rws.append(["Estado del parser", PARSER_STATUS])
        buff = io.BytesIO(); wb.save(buff); buff.seek(0)
        return StreamingResponse(buff, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'})

    raise HTTPException(404, "Formato no disponible")


@app.get("/api/project/export/corrections.json")
def export_corrections(documents: str | None = None):
    rows = []
    for doc in _project_docs(documents):
        for c in doc.get("corrections", []):
            rows.append({"document_id": doc["id"], "thesis_id": doc.get("thesis_id"), "thesis_title": doc.get("title"), "filename": doc["filename"], **c})
    return Response(json.dumps({"schema": 1, "corrections": rows}, ensure_ascii=False, indent=2), media_type="application/json", headers={"Content-Disposition": 'attachment; filename="biblioref_human_corrections.json"'})


GROUPINGS = ("university", "program", "faculty", "thesis_year", "director", "degree_level")


class AnalysisAliasPayload(BaseModel):
    project_id: str
    field: str
    variant: str
    canonical: str | None = None


class AnalysisRunPayload(BaseModel):
    group_by: str = "university"
    project_id: str | None = None


def _derived_connection():
    """Open the derived sidecar, creating it on first use.

    The sidecar lives beside project.json and is disposable: deleting it loses
    only the cached external answers, never extraction or manual work.
    """
    from biblioref.derived import open_store
    return open_store(WORKSPACE / "derived.sqlite3")


def _check_grouping(group_by: str) -> str:
    if group_by not in GROUPINGS:
        raise HTTPException(400, f"Agrupación no soportada: {group_by}")
    return group_by


@app.get("/api/analysis")
def get_analysis(group_by: str = "university", project_id: str | None = None):
    """The last analysis the user asked for, if any.

    Nothing is computed here. An empty result is the honest state of a project
    that has not been analysed yet, and the screen shows its button.
    """
    from biblioref.derived.analysis import stored_report
    from biblioref.derived.sync import source_fingerprint

    _check_grouping(group_by)
    scope = project_id or ""
    with STORE.lock:
        # Staleness is measured against the project store, not against the
        # sidecar: an edit made since the last run has not reached it yet.
        fingerprint = source_fingerprint(STORE.data, scope or None)
    conn = _derived_connection()
    try:
        stored = stored_report(conn, group_by=group_by, fingerprint=fingerprint, project_id=scope)
    finally:
        conn.close()
    if stored is None:
        return {"status": "never_run", "group_by": group_by, "project_id": scope, "report": None}
    return {"status": "stale" if stored["stale"] else "current", **stored}


@app.post("/api/analysis/run")
def run_analysis(payload: AnalysisRunPayload):
    """Synchronise the sidecar and compute the report. Offline, no network."""
    from biblioref.derived.analysis import run_and_store
    from biblioref.derived.sync import sync_documents

    from biblioref.derived.sync import documents_of, source_fingerprint

    group_by = _check_grouping(payload.group_by)
    scope = payload.project_id or ""
    conn = _derived_connection()
    try:
        with STORE.lock:
            sync = sync_documents(conn, STORE.data)
            doc_ids = documents_of(STORE.data, scope or None)
            fingerprint = source_fingerprint(STORE.data, scope or None)
        if scope and not doc_ids:
            raise HTTPException(404, "El proyecto no tiene tesis para analizar")
        result = run_and_store(
            conn, group_by=group_by, project_id=scope,
            doc_ids=doc_ids or None, fingerprint=fingerprint,
        )
    finally:
        conn.close()
    return {"status": "current", "sync": sync.as_dict(), **result}


@app.get("/api/analysis/alias-suggestions")
def analysis_alias_suggestions(group_by: str = "university", project_id: str | None = None):
    """Labels similar enough to ask whether they name the same thing.

    Nothing is merged here. Two labels that look alike are sometimes two
    different places, and only the user knows which.
    """
    from biblioref.derived.aliases import ALIASABLE_FIELDS, aliases_for, suggest_aliases
    from biblioref.derived.sync import documents_of

    _check_grouping(group_by)
    scope = project_id or ""
    if group_by not in ALIASABLE_FIELDS:
        return {"group_by": group_by, "aliasable": False, "suggestions": [], "aliases": {}}

    with STORE.lock:
        doc_ids = documents_of(STORE.data, scope or None)
        known = aliases_for(STORE.data, scope or None, group_by)
    conn = _derived_connection()
    try:
        suggestions = suggest_aliases(conn, group_by, doc_ids or None, known=known)
    finally:
        conn.close()
    return {"group_by": group_by, "aliasable": True, "suggestions": suggestions, "aliases": known}


@app.post("/api/analysis/aliases")
def set_analysis_alias(payload: AnalysisAliasPayload):
    """Record, or undo, a decision that two labels name the same thing.

    The decision lives in project.json beside the rest of the manual work: it
    is a human judgement, not derived data, so a rebuilt sidecar must not lose
    it. The original spelling in each document is never overwritten.
    """
    from biblioref.derived.aliases import ALIASABLE_FIELDS

    if payload.field not in ALIASABLE_FIELDS:
        raise HTTPException(400, f"Campo no unificable: {payload.field}")
    variant = (payload.variant or "").strip()
    canonical = (payload.canonical or "").strip()
    if not variant:
        raise HTTPException(400, "Falta la variante a unificar")
    if canonical == variant:
        raise HTTPException(400, "La variante y la forma canónica son la misma")

    with STORE.lock:
        projects = STORE.data.setdefault("projects", {})
        project = projects.get(payload.project_id)
        if not project:
            raise HTTPException(404, "Proyecto no encontrado")
        table = project.setdefault("group_aliases", {}).setdefault(payload.field, {})
        if canonical:
            if table.get(canonical) == variant:
                raise HTTPException(400, "Esa unificación invertiría una ya registrada")
            table[variant] = canonical
        else:
            table.pop(variant, None)
        if not table:
            project["group_aliases"].pop(payload.field, None)
        project["updated_at"] = _now()
        STORE.save()
        aliases = dict(project.get("group_aliases", {}).get(payload.field, {}))

    return {"field": payload.field, "aliases": aliases}


@app.get("/api/analysis/export.json")
def export_analysis(group_by: str = "university", project_id: str | None = None):
    from biblioref.derived.analysis import stored_report

    _check_grouping(group_by)
    conn = _derived_connection()
    try:
        stored = stored_report(conn, group_by=group_by, project_id=project_id or "")
    finally:
        conn.close()
    if stored is None:
        raise HTTPException(404, "Todavía no se analizó la bibliografía")
    return Response(
        json.dumps(stored, ensure_ascii=False, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": 'attachment; filename="biblioref_analisis.json"'},
    )


def run(port: int = 8765, open_browser: bool = True):
    import uvicorn
    if open_browser:
        threading.Timer(1.0, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="warning")


if __name__ == "__main__":
    run()
