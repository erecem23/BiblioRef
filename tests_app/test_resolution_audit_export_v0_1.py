import io
import zipfile

import pandas as pd
from openpyxl import load_workbook

from biblioref.audit.export import build_audit_bundle, workbook_bytes
from biblioref.audit.resolution import attach_resolution_columns, resolution_audit_columns
from biblioref.resolver.decision import ResolutionDecision
from biblioref.resolver.models import Candidate, CandidateBatch, ResolverQuery


def sample_sidecar():
    query = ResolverQuery(raw_reference="Cohen 1992 A power primer")
    candidate = Candidate(
        source="crossref",
        source_id="10.1037/0033-2909.112.1.155",
        doi="10.1037/0033-2909.112.1.155",
        title="A power primer",
        year="1992",
    )
    batch = CandidateBatch(query=query, candidates=(candidate,))
    decision = ResolutionDecision(
        version="resolution-decision-v0.1",
        status="VERIFIED_EXTERNAL_CONSENSUS",
        selected_candidate_index=0,
        selected_identity_key=candidate.identity_key,
        decision_method="cross_source_consensus",
        review_required=False,
        evidence_basis=("same_identity_across_independent_sources",),
        warnings=(),
        consensus_version="metadata-consensus-v0.1",
        consensus_status="external_consensus",
    )
    return batch, decision


def test_resolution_audit_columns_expose_selected_candidate_with_provenance():
    batch, decision = sample_sidecar()
    row = resolution_audit_columns(batch, decision)
    assert row["resolution_status"] == "VERIFIED_EXTERNAL_CONSENSUS"
    assert row["resolution_selected_doi"] == "10.1037/0033-2909.112.1.155"
    assert row["resolution_selected_source"] == "crossref"
    assert row["resolution_review_required"] is False


def test_attach_resolution_columns_requires_one_sidecar_per_record():
    batch, decision = sample_sidecar()
    rows = attach_resolution_columns([{"raw_reference": "x"}], [(batch, decision)])
    assert rows[0]["resolution_selected_title"] == "A power primer"


def test_workbook_adds_resolution_sheet_without_breaking_old_sheets():
    resolution = pd.DataFrame([resolution_audit_columns(*sample_sidecar())])
    data = workbook_bytes(
        pd.DataFrame([{"reference": "x"}]),
        pd.DataFrame([{"page": 1}]),
        pd.DataFrame(),
        resolution=resolution,
    )
    book = load_workbook(io.BytesIO(data), read_only=True)
    assert book.sheetnames == ["References", "Zones", "Manual clicks", "Resolution"]


def test_bundle_exports_resolution_csv_and_summary_without_rendering_screenshots():
    resolution = pd.DataFrame([resolution_audit_columns(*sample_sidecar())])
    data = build_audit_bundle(
        "unused.pdf",
        {},
        pd.DataFrame([{"accepted": True}]),
        pd.DataFrame(),
        pd.DataFrame(),
        (),
        resolution=resolution,
    )
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
        assert "resolution_decisions.csv" in names
        report = archive.read("processing_report_audited.json").decode("utf-8")
        assert '"n_resolution_rows": 1' in report
        assert '"VERIFIED_EXTERNAL_CONSENSUS": 1' in report
