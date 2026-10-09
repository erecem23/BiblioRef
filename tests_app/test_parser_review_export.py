# -*- coding: utf-8 -*-
from __future__ import annotations

import io
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from biblioref.parsing.review_export import _augment_workbook


class _Store:
    def docs(self):
        return {
            "doc1": {
                "id": "doc1",
                "thesis_id": "T0001",
                "title": "Tesis sintética",
                "references": [
                    {
                        "id": "ref1",
                        "order": 1,
                        "text": "Rivas, C. (2021). Nota. Retrieved 2022, from Portal: https://example.org/nota",
                        "parsed": {
                            "needs_review": True,
                            "review_reasons": "multiple_year_mentions",
                            "parser_notes": "year_primary_parenthetical_recovered",
                            "parser_confidence": 0.9,
                            "year": "2021",
                            "year_candidates": "2021 [publication] | 2022 [access]",
                            "authors": "Rivas, C.",
                            "author_candidates": "",
                            "title": "Nota",
                            "title_candidates": "",
                            "journal": None,
                            "journal_candidates": "",
                            "publication_type": "material_web",
                            "type_candidates": "",
                            "publisher": None,
                            "publisher_candidates": "",
                        },
                    }
                ],
            }
        }


def _visible_refs(doc):
    return list(doc.get("references") or [])


def test_augment_workbook_adds_review_columns_and_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Referencias"
    ws.append(["Tesis ID", "Título de tesis", "Referencia ID", "N", "Referencia"])
    ws.append(["T0001", "Tesis sintética", "ref1", 1, "Referencia sintética"])
    raw = io.BytesIO()
    wb.save(raw)

    app_module = SimpleNamespace(STORE=_Store(), _visible_refs=_visible_refs)
    out = _augment_workbook(raw.getvalue(), app_module)
    check = load_workbook(io.BytesIO(out), data_only=True)

    headers = [c.value for c in check["Referencias"][1]]
    assert "Revisar parser" in headers
    assert "Años candidatos" in headers
    review = check["Revisión parser"]
    assert review.max_row == 2
    assert review.cell(2, 5).value == "multiple_year_mentions"
    assert "2021 [publication]" in review.cell(2, 8).value
