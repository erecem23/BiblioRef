# -*- coding: utf-8 -*-
"""Las planillas que escribe BiblioRef sobreviven a un texto ilegible.

Una sola referencia con caracteres de control —OCR, o una capa de texto sin
forma de volver a las letras (T0140)— hacía fallar la exportación entera con
IllegalCharacterError: desde «Exportar» se veía como un botón que no hace nada.
"""
import os
import tempfile

os.environ.setdefault("BIBLIOREF_WORKSPACE", tempfile.mkdtemp())

import openpyxl

import app


def test_control_characters_are_stripped_from_every_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tesis", "D\x00?;E\x01<I?;JIK", 2025])
    ws.append(("tupla", "texto\x0bcortado"))
    assert ws.cell(1, 2).value == "D?;E<I?;JIK"
    assert ws.cell(2, 2).value == "textocortado"
    assert ws.cell(1, 3).value == 2025


def test_plain_text_and_accents_are_untouched():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cía, A.H. (2000). La ansiedad y sus trastornos — pp. 31-59."])
    assert ws.cell(1, 1).value == "Cía, A.H. (2000). La ansiedad y sus trastornos — pp. 31-59."
    assert app._xl(None) is None


def test_the_sanitiser_is_installed_once():
    app._install_xlsx_sanitiser()
    app._install_xlsx_sanitiser()
    wb = openpyxl.Workbook()
    wb.active.append(["a\x00b"])
    assert wb.active.cell(1, 1).value == "ab"
