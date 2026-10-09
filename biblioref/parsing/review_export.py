# -*- coding: utf-8 -*-
"""Export integration for parser uncertainty metadata.

This module is intentionally outside ``app.py`` so the parser audit layer can
be added without coupling it to extraction/UI internals.  It extends CSV field
selection and post-processes generated XLSX workbooks with parser review
columns plus a dedicated ``Revisión parser`` sheet.
"""
from __future__ import annotations

import io
from typing import Any, Dict

from fastapi.responses import Response
from openpyxl import load_workbook


REVIEW_FLAT_FIELDS = [
    "parsed_needs_review",
    "parsed_review_reasons",
    "parsed_parser_notes",
    "parsed_parser_confidence",
    "parsed_year_candidates",
    "parsed_author_candidates",
    "parsed_title_candidates",
    "parsed_journal_candidates",
    "parsed_type_candidates",
    "parsed_publisher_candidates",
]

REVIEW_XLSX_HEADERS = [
    "Revisar parser",
    "Motivos revisión parser",
    "Notas parser",
    "Confianza parser",
    "Años candidatos",
    "Autores candidatos",
    "Títulos candidatos",
    "Revistas candidatas",
    "Tipos candidatos",
    "Editoriales candidatas",
]


def _parser_values(parsed: Dict[str, Any]) -> list[Any]:
    return [
        bool(parsed.get("needs_review")),
        parsed.get("review_reasons"),
        parsed.get("parser_notes"),
        parsed.get("parser_confidence"),
        parsed.get("year_candidates"),
        parsed.get("author_candidates"),
        parsed.get("title_candidates"),
        parsed.get("journal_candidates"),
        parsed.get("type_candidates"),
        parsed.get("publisher_candidates"),
    ]


def _reference_map(app_module) -> Dict[str, tuple[dict, dict]]:
    out: Dict[str, tuple[dict, dict]] = {}
    for doc in app_module.STORE.docs().values():
        for ref in app_module._visible_refs(doc):
            rid = str(ref.get("id") or "")
            if rid:
                out[rid] = (doc, ref)
    return out


def _augment_workbook(body: bytes, app_module) -> bytes:
    wb = load_workbook(io.BytesIO(body))
    if "Referencias" not in wb.sheetnames:
        return body
    ws = wb["Referencias"]
    headers = [cell.value for cell in ws[1]]
    if "Referencia ID" not in headers:
        return body
    ref_col = headers.index("Referencia ID") + 1
    n_col = headers.index("N") + 1 if "N" in headers else None
    thesis_col = headers.index("Tesis ID") + 1 if "Tesis ID" in headers else None
    title_col = headers.index("Título de tesis") + 1 if "Título de tesis" in headers else None

    # Idempotent: a repeated middleware/install must not duplicate columns.
    start_col = ws.max_column + 1
    if REVIEW_XLSX_HEADERS[0] in headers:
        start_col = headers.index(REVIEW_XLSX_HEADERS[0]) + 1
    else:
        for offset, label in enumerate(REVIEW_XLSX_HEADERS):
            ws.cell(row=1, column=start_col + offset, value=label)

    refs = _reference_map(app_module)
    review_rows = []
    for row_idx in range(2, ws.max_row + 1):
        rid = str(ws.cell(row=row_idx, column=ref_col).value or "")
        pair = refs.get(rid)
        parsed = ((pair[1].get("parsed") or {}) if pair else {})
        values = _parser_values(parsed)
        for offset, value in enumerate(values):
            ws.cell(row=row_idx, column=start_col + offset, value=value)

        if parsed.get("needs_review"):
            doc = pair[0] if pair else {}
            ref = pair[1] if pair else {}
            review_rows.append([
                ws.cell(row=row_idx, column=thesis_col).value if thesis_col else doc.get("thesis_id"),
                ws.cell(row=row_idx, column=title_col).value if title_col else doc.get("title"),
                rid,
                ws.cell(row=row_idx, column=n_col).value if n_col else ref.get("order"),
                parsed.get("review_reasons"),
                parsed.get("parser_confidence"),
                parsed.get("year"),
                parsed.get("year_candidates"),
                parsed.get("authors"),
                parsed.get("author_candidates"),
                parsed.get("title"),
                parsed.get("title_candidates"),
                parsed.get("journal"),
                parsed.get("journal_candidates"),
                parsed.get("publication_type"),
                parsed.get("type_candidates"),
                parsed.get("publisher"),
                parsed.get("publisher_candidates"),
                ref.get("text"),
            ])

    # Human audit/corrections remain in the pre-existing ``Revisión`` sheet.
    # Parser uncertainty gets its own sheet so the two concepts never mix.
    if "Revisión parser" in wb.sheetnames:
        del wb["Revisión parser"]
    rws = wb.create_sheet("Revisión parser")
    rws.append([
        "Tesis ID", "Título de tesis", "Referencia ID", "N", "Motivos",
        "Confianza", "Año seleccionado", "Años candidatos", "Autores seleccionados",
        "Autores candidatos", "Título seleccionado", "Títulos candidatos",
        "Revista seleccionada", "Revistas candidatas", "Tipo seleccionado",
        "Tipos candidatos", "Editorial seleccionada", "Editoriales candidatas",
        "Referencia original",
    ])
    for row in review_rows:
        rws.append(row)
    rws.freeze_panes = "A2"
    rws.auto_filter.ref = rws.dimensions
    rws.column_dimensions["A"].width = 12
    rws.column_dimensions["B"].width = 45
    rws.column_dimensions["E"].width = 34
    rws.column_dimensions["H"].width = 42
    rws.column_dimensions["J"].width = 50
    rws.column_dimensions["L"].width = 60
    rws.column_dimensions["N"].width = 45
    rws.column_dimensions["P"].width = 35
    rws.column_dimensions["R"].width = 35
    rws.column_dimensions["S"].width = 100

    buff = io.BytesIO()
    wb.save(buff)
    return buff.getvalue()


def install_parser_review_export(app_module) -> None:
    """Install CSV/XLSX review export hooks exactly once on a BiblioRef app."""
    app = app_module.app
    if getattr(app.state, "parser_review_export_installed", False):
        return
    app.state.parser_review_export_installed = True

    for field in REVIEW_FLAT_FIELDS:
        if field not in app_module.PARSED_EXPORT_FIELDS:
            app_module.PARSED_EXPORT_FIELDS.append(field)

    @app.middleware("http")
    async def parser_review_xlsx_middleware(request, call_next):
        response = await call_next(request)
        path = request.url.path.rstrip("/").lower()
        ctype = (response.headers.get("content-type") or "").lower()
        if response.status_code != 200 or not path.endswith("/xlsx") or "spreadsheetml" not in ctype:
            return response

        chunks = []
        async for chunk in response.body_iterator:
            chunks.append(chunk)
        body = b"".join(chunks)
        try:
            body = _augment_workbook(body, app_module)
        except Exception:
            # Export must never fail merely because the audit decoration failed.
            pass

        headers = dict(response.headers)
        headers.pop("content-length", None)
        return Response(
            content=body,
            status_code=response.status_code,
            headers=headers,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
